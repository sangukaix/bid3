from io import BytesIO
from django.test import TestCase,override_settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from tempfile import TemporaryDirectory
from pptx import Presentation
from pptx.util import Inches,Pt
from hashlib import sha256
from unittest.mock import patch
from django.utils import timezone
from .models import BidNotice,SavedBid,BidProposal


class FinalizeReviewTests(TestCase):
    def setUp(self):
        self.directory=TemporaryDirectory();self.addCleanup(self.directory.cleanup)
        self.override=override_settings(MEDIA_ROOT=self.directory.name);self.override.enable();self.addCleanup(self.override.disable)
        self.user=get_user_model().objects.create_user(username='review-owner')
        notice=BidNotice.objects.create(bid_ntce_no='REVIEW-001',title='교육 운영',business_type='용역')
        saved=SavedBid.objects.create(user=self.user,bid_notice=notice)
        self.proposal=BidProposal.objects.create(saved_bid=saved,output_format='pptx',revision_plan={'status':'draft','quality_review':{'passed':True}})
        self.client=APIClient();self.client.force_authenticate(self.user)

    def upload(self,text,small=False):
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        box=page.shapes.add_textbox(0,0,Inches(3 if small else 8),Inches(.3 if small else 1))
        box.text=text;box.text_frame.paragraphs[0].runs[0].font.size=Pt(36 if small else 20)
        buffer=BytesIO();prs.save(buffer)
        self.proposal.generated_file.save('review.pptx',SimpleUploadedFile('review.pptx',buffer.getvalue()))

    def finalize(self,data=None):
        return self.client.post('/api/bids/REVIEW-001/proposal/finalize/',data or {},format='json')

    def test_stale_pass_does_not_allow_template_text_even_with_layout_acknowledgment(self):
        self.upload('사업명 제안서')
        response=self.finalize({'acknowledge_layout_warnings':True})
        self.assertEqual(response.status_code,409)
        self.proposal.refresh_from_db();self.assertEqual(self.proposal.revision_plan['status'],'draft')
        self.assertFalse(response.data['proposal']['revision_plan']['quality_review']['passed'])

    def test_severe_estimated_overflow_requires_explicit_visual_confirmation(self):
        self.upload('지역 영어 교육 운영과 결과 보고를 위한 수행 계획입니다. '*5,small=True)
        original=self.proposal.generated_file.read();self.proposal.generated_file.close()
        self.assertEqual(self.finalize().status_code,409)
        self.assertEqual(self.finalize({'acknowledge_layout_warnings':'true'}).status_code,409)
        self.assertEqual(self.finalize({'acknowledge_layout_warnings':True}).status_code,200)
        self.proposal.refresh_from_db();self.assertEqual(self.proposal.generated_file.read(),original)
        self.proposal.generated_file.close()

    def test_invalid_package_is_not_marked_final(self):
        self.proposal.generated_file.save('bad.pptx',SimpleUploadedFile('bad.pptx',b'broken'))
        self.assertEqual(self.finalize().status_code,422)
        self.proposal.refresh_from_db();self.assertEqual(self.proposal.revision_plan['status'],'draft')

    def output_review(self):
        return self.client.post('/api/bids/REVIEW-001/proposal/output-review/',{},format='json')

    def test_review_invalidates_stale_coverage_without_changing_the_file(self):
        self.upload('실제 운영 방안을 설명합니다.')
        original=self.proposal.generated_file.read();self.proposal.generated_file.close()
        self.proposal.revision_plan.update({
            'requirement_register':{'requirements':[{'requirement':'250명에게 30분 수업'}]},
            'requirement_coverage':{'covered_count':1,'checks':[{'id':'R0001',
                'requirement':'250명에게 30분 수업','covered':True,'slide_number':1,
                'quote':'250명에게 30분 수업을 운영합니다.'}]}})
        self.proposal.save()
        response=self.output_review()
        self.assertEqual(response.status_code,200)
        self.proposal.refresh_from_db()
        self.assertEqual(self.proposal.revision_plan['requirement_coverage']['covered_count'],0)
        self.assertEqual(self.proposal.revision_plan['output_review']['file_sha256'],sha256(original).hexdigest())
        self.assertEqual(self.proposal.generated_file.read(),original);self.proposal.generated_file.close()
        self.assertEqual(self.proposal.revision_plan['status'],'draft')

    def test_other_user_cannot_review_and_invalid_file_is_rejected(self):
        self.upload('수행 계획을 설명합니다.')
        self.client.force_authenticate(get_user_model().objects.create_user(username='other-reviewer'))
        self.assertEqual(self.output_review().status_code,404)
        self.client.force_authenticate(self.user)
        self.proposal.generated_file.save('bad.pptx',SimpleUploadedFile('bad.pptx',b'broken'))
        self.assertEqual(self.output_review().status_code,422)

    def test_concurrent_finalization_is_preserved_instead_of_overwritten(self):
        self.upload('수행 계획을 설명합니다.')
        from .services.proposal_output_review import refresh_output_review
        def concurrent_change(content,plan):
            BidProposal.objects.filter(pk=self.proposal.pk).update(
                revision_plan={'status':'final'},updated_at=timezone.now())
            return refresh_output_review(content,plan)
        with patch('bids.services.proposal_output_review.refresh_output_review',side_effect=concurrent_change):
            self.assertEqual(self.output_review().status_code,409)
        self.proposal.refresh_from_db();self.assertEqual(self.proposal.revision_plan['status'],'final')

    def test_finalize_also_records_actual_file_audit(self):
        self.upload('수행 계획을 설명합니다.')
        self.assertEqual(self.finalize().status_code,200)
        self.proposal.refresh_from_db()
        self.assertFalse(self.proposal.revision_plan['output_review']['source_register_available'])
        self.assertEqual(self.proposal.revision_plan['output_review']['actual_slide_count'],1)
