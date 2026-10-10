"""Read-back audit of the actual PPTX. Never interprets a match as bid eligibility."""
from hashlib import sha256
from io import BytesIO
import re

from pptx import Presentation
from .proposal_coverage import CoverageVerdict, confirmed_quote, normalize
from .proposal_evidence import requirement_rows

VERSION = 'proposal-output-evidence-v1'
UNFILLED = re.compile(r'확인\s*필요|자료\s*미제공|자료\s*미확인|증빙\s*미제공|작성\s*예정|미정')


def visible_text_blocks(content, exclude_conditions=False, exclude_labels=False):
    """Read complete text, including fixed shapes, groups and tables; exclude notes."""
    prs = Presentation(BytesIO(content))
    def blocks(shapes, prefix=''):
        for index, shape in enumerate(shapes):
            target = f'{prefix}shape-{index}'
            if exclude_conditions and shape.name.startswith('bid3-requirement-'):
                continue
            if exclude_labels and (shape.name in {'bid3-title','cover-title','section-title'}
                    or shape.name.startswith(('footer-','section-label-','number-','title-'))):
                continue
            if shape.shape_type == 6:
                yield from blocks(shape.shapes, target+'/')
            elif getattr(shape, 'has_table', False):
                for row_index, row in enumerate(shape.table.rows):
                    for cell_index, cell in enumerate(row.cells):
                        if not cell.is_spanned and cell.text.strip():
                            yield {'target':f'{target}-cell-{row_index}-{cell_index}','text':cell.text}
            elif getattr(shape, 'has_text_frame', False) and shape.text.strip():
                yield {'target':target,'name':shape.name,'text':shape.text}
    return {number:list(blocks(slide.shapes)) for number,slide in enumerate(prs.slides,1)}


def written_pages(content, exclude_conditions=False, exclude_labels=False):
    return {number:'\n'.join(block['text'] for block in blocks)
            for number,blocks in visible_text_blocks(content,exclude_conditions,exclude_labels).items()}


def output_plan(content):
    """The semantic reviewer receives the exported pages, not proposed changes."""
    return {'slide_changes': [{'slide_number':number, 'action':'UPDATE',
        'text_changes':[{'revised_text':block['text'],'shape_name':block.get('name','')} for block in blocks]}
        for number,blocks in visible_text_blocks(content,exclude_conditions=True).items()]}


def audit_output(content, plan):
    pages = written_pages(content,exclude_conditions=True)
    rows = requirement_rows(plan.get('requirement_register', {}))
    previous = {item.get('id'):item for item in (plan.get('requirement_coverage') or {}).get('checks', [])}
    checks = []
    for row in rows:
        old = previous.get(row['id'], {})
        quote = normalize(old.get('quote', ''))
        same_requirement = normalize(old.get('requirement', '')) == normalize(row.get('requirement', ''))
        matches = [n for n,text in pages.items() if len(quote)>=8 and quote in normalize(text)]
        # A former page may have moved after insertions. Rebind only unambiguously.
        page = old.get('slide_number') if old.get('slide_number') in matches else (matches[0] if len(matches)==1 else None)
        verdict = CoverageVerdict(covered=old.get('covered') is True and same_requirement,
                                 slide_number=page or 0, quote=old.get('quote', '')[:600])
        matched = confirmed_quote(verdict, row.get('requirement',''), pages)
        checks.append({**row, 'covered':bool(matched), 'slide_number':page if matched else None,
            'quote':old.get('quote','') if matched else '',
            'status':'passage_matched' if matched else 'unverified' if old.get('status')=='unverified' else 'review_required',
            'reason':old.get('reason','') if matched else (
                '이전 검수 인용이 현재 파일의 요구 조건과 일치하지 않습니다.' if old.get('covered')
                else old.get('reason') or '현재 파일에서 해당 요구사항의 답변을 검수해야 합니다.')})
    open_text = [{'slide_number':number, 'text':line.strip()} for number,text in written_pages(content).items()
                 for line in text.splitlines() if UNFILLED.search(line)]
    notes = []
    if not rows:
        notes.append('공고 요구사항 목록이 저장되지 않아 전체 반영 여부는 미검수입니다. 원문을 확인하고 새 생성 시 목록을 함께 저장하세요.')
    if checks and any(not item['covered'] for item in checks):
        notes.append('본문 답변을 확인하지 못한 요구사항이 있습니다. 별도 제출서류·자격 요건도 원문과 증빙으로 확인하세요.')
    failure_count=sum(item['status']=='unverified' for item in checks)
    if failure_count:
        notes.append(f'자동 대조 실패 {failure_count}개가 남아 있습니다. 전체 AI 검수를 다시 실행하거나 원문을 직접 대조하세요.')
    processing = plan.get('document_processing') or {}
    failed = processing.get('failed_files') or []
    if failed:
        notes.append(f'공고 첨부파일 {len(failed)}개를 읽지 못했습니다. 원문 누락 여부를 확인하세요.')
    if open_text:
        notes.append('확인 필요·미정·작성 예정 문구가 실제 PPTX에 남아 있습니다.')
    limit = plan.get('detected_page_limit')
    limit_exceeded = isinstance(limit,int) and not isinstance(limit,bool) and limit>0 and len(pages)>limit
    if limit_exceeded:
        notes.append(f'저장된 공고 분량 제한 {limit}쪽보다 출력 {len(pages)}쪽이 많습니다.')
    return {'version':VERSION, 'file_sha256':sha256(content).hexdigest(),
        'scope':'최종 PPTX 본문의 요구사항 인용 대조', 'source_register_available':bool(rows),
        'actual_slide_count':len(pages), 'matched_count':sum(item['covered'] for item in checks),
        'total_count':len(rows), 'review_required_count':sum(not item['covered'] for item in checks),
        'verification_failure_count':failure_count,
        'checks':checks, 'open_text_items':open_text, 'failed_files':failed,
        'page_limit_exceeded':limit_exceeded, 'review_notes':notes,
        'limitation':'본문 인용 일치는 자격·증빙의 진위, 조건 전체 충족, 평가 점수 또는 제출 가능성을 보증하지 않습니다.'}


def refresh_output_review(content, plan):
    """Invalidate stale coverage counters together with their old passages."""
    from .proposal_final_review import invalidate_final_review
    invalidate_final_review(content, plan)
    report = audit_output(content, plan)
    plan['output_review'] = report
    plan['requirement_coverage'] = {**(plan.get('requirement_coverage') or {}),
        'covered_count':report['matched_count'], 'total_count':report['total_count'],
        'checks':report['checks'], 'missing_requirements':[i['requirement'] for i in report['checks'] if not i['covered']],
        'unverified_count':report['verification_failure_count'], 'reviewed_artifact':'exported_pptx',
        'source_register_available':report['source_register_available'], 'file_sha256':report['file_sha256']}
    return report
