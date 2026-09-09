from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


KOREAN_FONT = "Malgun Gothic"

def _set_font(element, font):
    """Word가 한글에 사용하는 eastAsia 글꼴까지 명시합니다."""

    properties = element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for font_type in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{font_type}"), font)
    fonts.set(qn("w:hint"), "eastAsia")


def _apply_korean_font(document):
    """제목·목록·표를 포함한 문서 전체에 한글 글꼴을 적용합니다."""

    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = KOREAN_FONT
            _set_font(style.element, KOREAN_FONT)

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = KOREAN_FONT
            _set_font(run._element, KOREAN_FONT)


def _add_bullets(document, values, empty_text="해당 없음"):
    for value in values or [empty_text]:
        document.add_paragraph(str(value), style="List Bullet")


def _add_form_table(document, form):
    """찾아낸 정량평가 양식을 작성값과 공란이 보이도록 표로 만듭니다."""

    document.add_heading(form.get("form_name") or "정량평가 양식", level=2)
    source = " · ".join(
        value
        for value in (form.get("source_file"), form.get("source_location"))
        if value
    )
    if source:
        document.add_paragraph(f"원본 위치: {source}")
    if form.get("purpose"):
        document.add_paragraph(form["purpose"])

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, value in zip(
        table.rows[0].cells,
        ("작성 항목", "작성 내용", "상태", "확인사항"),
    ):
        cell.text = value

    for field in form.get("fields", []):
        cells = table.add_row().cells
        cells[0].text = field.get("label") or "항목명 확인 필요"
        cells[1].text = str(field.get("value") or "")  # 근거가 없으면 공란 유지
        cells[2].text = field.get("status") or "미작성"
        cells[3].text = field.get("missing_reason") or ""


def build_quantitative_docx(bid_notice, company_name, report):
    """원본에서 찾은 정량평가 양식을 채운 Word 검토본을 만듭니다."""

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = KOREAN_FONT
    normal.font.size = Pt(9.5)

    title = document.add_heading("정량평가 양식 작성본", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"공고명: {bid_notice.title}")
    document.add_paragraph(f"회사명: {company_name}")
    document.add_paragraph(
        "공고 첨부 양식에서 확인한 항목 중 회사 자료로 확인되는 값만 작성했습니다. "
        "근거가 없는 항목은 공란으로 남겼습니다."
    )

    completion = report.get("completion", {})
    document.add_heading("1. 작성 현황", level=1)
    summary_table = document.add_table(rows=2, cols=4)
    summary_table.style = "Table Grid"
    labels = ("확인된 양식", "전체 항목", "작성 완료", "남은 항목")
    values = (
        completion.get("form_count", len(report.get("forms", []))),
        completion.get("total_fields", 0),
        completion.get("completed_fields", 0),
        completion.get("incomplete_fields", 0)
        + completion.get("direct_review_fields", 0),
    )
    for cell, value in zip(summary_table.rows[0].cells, labels):
        cell.text = value
    for cell, value in zip(summary_table.rows[1].cells, values):
        cell.text = str(value)

    document.add_heading("2. 양식별 작성 내용", level=1)
    forms = report.get("forms", [])
    if forms:
        for form in forms:
            _add_form_table(document, form)
    else:
        document.add_paragraph("첨부문서에서 별도 작성 양식을 확인하지 못했습니다.")

    document.add_heading("3. 남은 작성 항목", level=1)
    remaining = (
        report.get("incomplete_field_names", [])
        + report.get("direct_review_field_names", [])
    )
    _add_bullets(document, remaining, empty_text="남은 항목 없음")

    document.add_heading("4. 추가 확보 자료", level=1)
    _add_bullets(document, report.get("missing_documents", []))

    document.add_heading("5. 최종 확인사항", level=1)
    _add_bullets(document, report.get("review_notes", []))

    document.add_heading("6. 출처", level=1)
    for source in report.get("sources", []):
        document.add_paragraph(
            f"[출처 {source.get('number')}] {source.get('file_name')} · {source.get('location')}",
            style="List Bullet",
        )

    _apply_korean_font(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
