"""Default-template cleanup and read-back compatibility for older detail pages."""
from io import BytesIO
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from .proposal_pptx_renderer import inspect_proposal_quality, _replace_text_frame, _slide_elements
from .text_geometry import measure_text, fitting_size
from .proposal_evidence import requirement_rows

FONT = '맑은 고딕'
BODY_SIZE = 12
INK = '172C48'
ACCENT = '147D78'


def _height(text, width, size=BODY_SIZE):
    measured = measure_text(text, {'width':width, 'height':10000, 'font_size':size,
        'font_family':FONT, 'margin_left':0, 'margin_right':0,
        'margin_top':0, 'margin_bottom':0, 'line_height_ratio':1.25})
    return measured['required_height'] + 4


def _text(slide, name, text, x, y, width, height, size=BODY_SIZE, color=INK, bold=False):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.text = text
    for paragraph in frame.paragraphs:
        # Use the same point leading as the height calculation. Office's
        # proportional leading includes font metrics and can overrun the box.
        paragraph.line_spacing = Pt(size * 1.25)
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)


def prepare_default_template_output(file_result, register, bid_notice=None):
    """Clean known illustrative copy and schedules only in fresh built-in-template output."""
    prs = Presentation(BytesIO(file_result['file_bytes']))
    rows = [row for row in requirement_rows(register)
            if re.search(r'^(?:교육|사업|용역|과업|수행|계약|운영)\s*기간\s*(?:[:：]|은)',row.get('requirement',''))]
    changed = []
    if bid_notice is not None:
        # The official title is a source fact, not copy for the model to abbreviate.
        cover = prs.slides[0]
        title = next((shape for shape in cover.shapes if shape.name in {'bid3-title','cover-title'}
                      and getattr(shape,'has_text_frame',False)),None)
        if title is not None and title.text != bid_notice.title:
            title.text_frame.word_wrap = True
            element = next(item for item in _slide_elements(cover) if item.get('name') == title.name)
            element = {**element,'width':element['width']*.96}
            size, _ = fitting_size(bid_notice.title, element, minimum=18)
            if size is None:
                raise ValueError('정식 공고명이 표지 제목 상자에 들어가지 않습니다. 더 넓은 표지 양식을 선택하세요.')
            lines = measure_text(bid_notice.title,element,size,include_lines=True)['lines']
            _replace_text_frame(title.text_frame,'\v'.join(lines),size)
            changed.append(1)
        for shape in cover.shapes:
            if shape.name not in {'domain-copy','cover-subtitle'} or not getattr(shape,'has_text_frame',False) or not shape.text.strip():
                continue
            element = next(item for item in _slide_elements(cover) if item.get('name')==shape.name)
            text = re.sub(r'(?<!\d)(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})(?!\d)',r'\1.\2.\3',shape.text)
            element = {**element,'width':element['width']*.96,'wrap':True}
            size,_ = fitting_size(text,element)
            if size is not None:
                lines = measure_text(text,element,size,include_lines=True)['lines']
                _replace_text_frame(shape.text_frame,'\v'.join(lines),size)
    neutral_copy = {
        '모든 제안 내용은 실제 보유한 실적, 인력, 인증 및 정량적 수치에 근거합니다.':'회사 실적·인력·인증을 검토 완료된 자료와 대조합니다.',
        '실제 수행 경험과 검증된 전문 인력이 사업의 성공을 보장합니다.':'회사 근거 확인 필요\n실적·인력·인증 자료를\n검토 후 반영합니다.',
        '사업과 직접 관련된 자료만 배치합니다.':'해당 사업의 회사 증빙 확인 필요',
    }
    for number,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if bid_notice is not None and getattr(shape,'has_text_frame',False):
                if shape.name=='bid3-fixed-design-name':
                    _replace_text_frame(shape.text_frame,'공고 '+getattr(bid_notice,'bid_ntce_no',''))
                elif shape.name=='bid3-fixed-status' and shape.text.strip()=='분야별 작성 양식':
                    _replace_text_frame(shape.text_frame,getattr(bid_notice,'notice_organization','') or '사업 수행 제안')
            if getattr(shape,'has_text_frame',False) and ' '.join(shape.text.split()) in neutral_copy:
                _replace_text_frame(shape.text_frame,neutral_copy[' '.join(shape.text.split())])
            if shape.name=='closing-title' and getattr(shape,'has_text_frame',False) and shape.text.strip()=='확인 필요':
                _replace_text_frame(shape.text_frame,'성과 목표와 검증 계획')
            if shape.name in {'understanding-source','closing-eyebrow'} and getattr(shape,'has_text_frame',False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size and run.font.size.pt < 11:
                            run.font.size = Pt(11)
            if shape.name.startswith('phase-output-') and getattr(shape,'has_text_frame',False):
                match = re.fullmatch(r'주요 산출물\s*\[([^\[\]]+)\]\s*',shape.text)
                if match:
                    _replace_text_frame(shape.text_frame,'주요 산출물\n'+match.group(1))
        names = {shape.name for shape in slide.shapes}
        if not {'schedule-header','week-0','schedule-bar-0'} <= names:
            continue
        # The built-in sample's fixed durations are not grounded in this RFP.
        for shape in list(slide.shapes):
            if shape.name.startswith(('schedule-','week-')):
                shape._element.getparent().remove(shape._element)
        width = prs.slide_width.pt - 108
        y, bottom = 160, prs.slide_height.pt - 70
        _text(slide,'schedule-source-heading','공고에 명시된 일정',54,y,width,23,16,ACCENT,True)
        y += 32
        for row in rows:
            text = row['requirement']
            height = _height(text,width)
            if y+height > bottom-45:
                break
            _text(slide,'bid3-requirement-schedule-'+row['id'],text,54,y,width,height)
            y += height+10
        message = '단계별 세부 기간은 착수 시 협의하여 확정합니다.'
        if not rows:
            message = '공고에서 사업·교육 기간을 확인하지 못했습니다. 단계별 세부 기간은 원문·담당자 확인 후 확정합니다.'
        _text(slide,'schedule-pending',message,54,min(y+8,bottom-40),width,40,12,'9A5B15')
        slide.notes_slide.notes_text_frame.text += '\n공고 일정 원문:\n'+'\n'.join(
            row['id']+' '+row['requirement']+' | '+' · '.join(row.get('sources',[])) for row in rows)
        changed.append(number)
    buffer=BytesIO();prs.save(buffer)
    result = {**file_result,'file_bytes':buffer.getvalue()}
    if changed:
        result['revision_log'] = [*file_result['revision_log'], *[
            {'source_slide_number':next((source for source,output in file_result.get('source_page_map',{}).items() if output==number),None),
             'output_slide_number':number,'action':'UPDATE',
             'title':'공고 원문 반영','reason':'정식 공고명·일정과 기본 안내를 공고 원문 기준으로 보완',
             'changes':[],'warnings':[]} for number in changed]]
    result['quality_review'] = inspect_proposal_quality(result['file_bytes'],result['revision_log'])
    return result


def bind_existing_details(content, writing_plan):
    """Reconnect a read-only full review to actual named answer boxes, not old page numbers."""
    rows = {row['id']:row for row in writing_plan.get('items',[]) if row.get('channel')=='body'}
    placements = {}
    def shapes_in(shapes):
        for shape in shapes:
            if shape.shape_type == 6:
                yield from shapes_in(shape.shapes)
            else:
                yield shape
    for number,slide in enumerate(Presentation(BytesIO(content)).slides,1):
        for shape in shapes_in(slide.shapes):
            if not shape.name.startswith('bid3-answer-') or not getattr(shape,'has_text_frame',False) or not shape.text.strip():
                continue
            key = shape.name.removeprefix('bid3-answer-')
            if key in rows:
                placements.setdefault(key,[]).append(number)
    if not placements:
        return
    for key,numbers in placements.items():
        rows[key]['detail_slide_numbers'] = sorted(set(numbers))
        rows[key]['output_slide_numbers'] = list(dict.fromkeys([*numbers,*rows[key].get('output_slide_numbers',[])]))
    writing_plan['detail_pages'] = {'added_slide_count':len({n for numbers in placements.values() for n in numbers}),
        'placed_count':len(placements),'body_count':len(rows), 'placements':{key:min(numbers) for key,numbers in placements.items()},
        'omitted':[{'id':key,'reason':'현재 파일에 전용 상세 답변 상자가 없습니다. 기존 본문 답변을 별도 대조하세요.'} for key in rows if key not in placements],
        'source':'actual_pptx'}
    writing_plan.setdefault('notes',[]).append(f'현재 파일에서 본문 {len(rows)}개 중 {len(placements)}개의 상세 답변 위치를 다시 연결했습니다. 위치 연결은 조건 충족 판정이 아닙니다.')
