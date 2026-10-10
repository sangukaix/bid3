from io import BytesIO
from django.test import SimpleTestCase
from pptx import Presentation
from pptx.util import Inches,Pt
from .services.proposal_pptx_renderer import inspect_proposal_quality, _slide_elements


class OutputReviewTests(SimpleTestCase):
    def test_total_period_and_consecutive_week_axis_are_reported_for_manual_review(self):
        for total,axis,expected in ((15,[1,2,3],1),(3,[1,2,3],0),(15,[1,3,5],0),(15,[1,2,999999999],0)):
            prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
            for i,text in enumerate([f'총 {total}주간의 교육 일정',*[f'{week}주' for week in axis]]):
                page.shapes.add_textbox(0,Inches(i),Inches(8),Inches(.5)).text=text
            buffer=BytesIO();prs.save(buffer);quality=inspect_proposal_quality(buffer.getvalue())
            self.assertEqual(len(quality['timeline_range_items']),expected)
            if expected:
                self.assertTrue(any('일정 축 범위 확인' in note for note in quality['review_items']))

    def test_tight_small_footer_can_be_edited_without_enlarging_or_wrapping(self):
        from .services.proposal_pptx_renderer import _apply_text_changes
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        caption=page.shapes.add_textbox(0,Inches(5),Inches(2),Inches(.2));caption.name='footer-title-1'
        caption.text='사업명 제안서';caption.text_frame.paragraphs[0].runs[0].font.size=Pt(8)
        applied,warnings=_apply_text_changes(page,[{'target':'shape-0','original_text':caption.text,'revised_text':'입찰 제안서'}])
        self.assertEqual(caption.text,'입찰 제안서');self.assertEqual(warnings,[])
        self.assertEqual(applied[0]['font_size_pt'],8)
        applied,warnings=_apply_text_changes(page,[{'target':'shape-0','revised_text':'줄바꿈이 필요한 매우 긴 설명입니다. '*20}])
        self.assertEqual(applied,[]);self.assertTrue(warnings);self.assertEqual(caption.text,'입찰 제안서')

    def test_one_line_with_large_font_is_not_exempt_just_because_named_as_footer(self):
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        caption=page.shapes.add_textbox(0,Inches(5),Inches(3),Inches(.2));caption.name='footer-page-1'
        caption.text='01';caption.text_frame.paragraphs[0].runs[0].font.size=Pt(40)
        buffer=BytesIO();prs.save(buffer);result=inspect_proposal_quality(buffer.getvalue())
        self.assertEqual(len(result['severe_overflow_items']),1)

    def test_tight_single_line_page_label_is_not_reported_as_small_body_or_clipped(self):
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        caption=page.shapes.add_textbox(0,Inches(5),Inches(2),Inches(.2));caption.name='footer-page-1'
        caption.text='01';caption.text_frame.paragraphs[0].runs[0].font.size=Pt(8)
        buffer=BytesIO();prs.save(buffer);result=inspect_proposal_quality(buffer.getvalue())
        self.assertEqual(result['small_text_items'],[]);self.assertEqual(result['overflow_items'],[])
    def test_plain_template_guidance_is_flagged_but_real_proposal_text_is_allowed(self):
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        for i,text in enumerate(('사업명 제안서','회사 이미지 · 수행 실적 · 인증 자료','정보시스템 구축 제안서')):
            page.shapes.add_textbox(0,Inches(i),Inches(8),Inches(1)).text=text
        buffer=BytesIO();prs.save(buffer);quality=inspect_proposal_quality(buffer.getvalue())
        self.assertEqual([i['marker'] for i in quality['template_leftovers']],['사업명 제안서','회사 이미지 · 수행 실적 · 인증 자료'])
        self.assertFalse(quality['passed'])

    def test_cover_title_is_recognized_and_severe_overflow_is_not_a_small_caption(self):
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        title=page.shapes.add_textbox(0,0,Inches(3),Inches(.5));title.name='cover-title'
        title.text='집중형 외국어 교육 프로그램 운영 용역 제안서'
        title.text_frame.paragraphs[0].runs[0].font.size=Pt(40)
        buffer=BytesIO();prs.save(buffer);quality=inspect_proposal_quality(buffer.getvalue())
        self.assertEqual(_slide_elements(page)[0]['kind'],'title')
        self.assertEqual(len(quality['severe_overflow_items']),1)
