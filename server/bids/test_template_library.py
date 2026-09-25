import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import tempfile

from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from pptx import Presentation

from .services.proposal_catalog import bundled_previews, load_studio_templates, template_slide_count
from .services.proposal_pptx_renderer import extract_pptx_inventory, build_proposal_pptx
from .services.proposal_preview import create_template_slide_previews


class TemplateLibraryTests(SimpleTestCase):
    def test_all_30_templates_have_editable_content_and_exact_previews(self):
        studio = load_studio_templates(settings.BASE_DIR / "proposal_templates/library")
        self.assertEqual(len(studio), 30)
        self.assertEqual(len({t["category"] for t in studio.values()}), 6)
        for template_id, template in studio.items():
            with self.subTest(template=template_id):
                self.assertEqual(template_slide_count(template["path"]), 20)
                self.assertEqual(len(bundled_previews(template)), 20)
                inventory = extract_pptx_inventory(template["path"])
                self.assertTrue(all(slide["elements"] for slide in inventory))
                self.assertEqual(inventory[10]["role"], "timeline")
                self.assertEqual(inventory[16]["role"], "metrics")
                self.assertTrue(any(e["kind"] == "table_cell" for e in inventory[5]["elements"]))
                self.assertTrue(any(e["kind"] == "title" for e in inventory[0]["elements"]))

    def test_registered_previews_need_no_office_conversion(self):
        with patch("bids.services.proposal_preview._convert_with_powerpoint", side_effect=AssertionError("unexpected conversion")):
            self.assertEqual(len(create_template_slide_previews("civic_navy")), 20)

    def test_stale_or_incomplete_bundled_previews_are_not_used(self):
        template = dict(settings.PROPOSAL_TEMPLATES["civic_navy"])
        with tempfile.TemporaryDirectory() as temporary:
            template["preview_directory"] = Path(temporary)
            (Path(temporary) / "manifest.json").write_text(json.dumps({"source_sha256":"wrong","slide_count":20}),encoding="utf-8")
            self.assertEqual(bundled_previews(template), [])
            original_manifest = Path(settings.PROPOSAL_TEMPLATES["civic_navy"]["preview_directory"]) / "manifest.json"
            (Path(temporary) / "manifest.json").write_bytes(original_manifest.read_bytes())
            self.assertEqual(bundled_previews(template), [])

    def test_catalog_rejects_traversal_and_duplicate_ids(self):
        for entries in ([{"id":"../secret","file":"../secret.pptx"}], [{"id":"okay","file":"../secret.pptx"}], [{"id":"okay","file":"okay.pptx"}]*2):
            with tempfile.TemporaryDirectory() as temporary:
                (Path(temporary) / "catalog.json").write_text(json.dumps(entries),encoding="utf-8")
                with self.assertRaises(ValueError):
                    load_studio_templates(temporary)

    def test_generation_preserves_images_and_edits_native_cells(self):
        for template_id in ("civic_navy", "municipal_teal", "classroom_apricot"):
            with self.subTest(template=template_id):
                source = settings.PROPOSAL_TEMPLATES[template_id]["path"]
                original = Presentation(source)
                inventory = extract_pptx_inventory(source)
                title = next(e for e in inventory[0]["elements"] if e["kind"] == "title")
                cell = next(e for e in inventory[5]["elements"] if e["kind"] == "table_cell" and "확인 필요" in e["text"])
                changes = [{"slide_number":n,"action":"UPDATE","reason":"검증","text_changes":[{"target":e["target"],"original_text":e["text"],"revised_text":value}]} for n,e,value in ((1,title,"지역 교육 운영 제안"),(6,cell,"대상 250명"))]
                # A generated plan must not be able to rewrite fixed page numbers.
                fixed_index = next(i for i,s in enumerate(original.slides[5].shapes) if s.name == "bid3-fixed-page")
                changes[1]["text_changes"].append({"target":f"shape-{fixed_index}","revised_text":"999"})
                result = build_proposal_pptx(source,SimpleNamespace(bid_ntce_no="TEMPLATE-TEST"),{"slide_changes":changes})
                result_deck = Presentation(BytesIO(result["file_bytes"]))
                self.assertEqual(len(result_deck.slides),20)
                self.assertIn("지역 교육 운영 제안", " ".join(s.text for s in result_deck.slides[0].shapes if s.has_text_frame))
                table = next(s.table for s in result_deck.slides[5].shapes if s.has_table)
                self.assertIn("대상 250명",[c.text for row in table.rows for c in row.cells])
                self.assertEqual(result_deck.slides[5].shapes[fixed_index].text,"06")
                old_images=[s.image.blob for s in original.slides[0].shapes if s.shape_type == 13]
                new_images=[s.image.blob for s in result_deck.slides[0].shapes if s.shape_type == 13]
                self.assertEqual(new_images,old_images)


class TemplateCatalogApiTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="template-library-test",password="test-only")

    def test_catalog_preview_and_download_require_authentication(self):
        for url in ("/api/proposal-templates/","/api/proposal-templates/civic_navy/download/","/api/proposal-templates/civic_navy/slides/1/"):
            self.assertIn(self.client.get(url).status_code,(401,403))

    def test_catalog_download_and_missing_template(self):
        self.client.force_login(self.user)
        response = self.client.get("/api/proposal-templates/")
        self.assertEqual(response.status_code,200)
        self.assertEqual(len(response.json()["templates"]),35)
        self.assertNotIn("path",response.json()["templates"][5])
        response = self.client.get("/api/proposal-templates/civic_navy/download/")
        self.assertEqual(response.status_code,200)
        self.assertIn("attachment",response["Content-Disposition"])
        self.assertTrue(b"".join(response.streaming_content).startswith(b"PK"))
        self.assertEqual(self.client.get("/api/proposal-templates/not-a-template/download/").status_code,404)
