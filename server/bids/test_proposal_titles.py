from io import BytesIO
from django.test import SimpleTestCase
from pptx import Presentation
from pptx.util import Inches, Pt
from bids.services.proposal_pptx_renderer import _slide_elements, _slide_title, inspect_proposal_quality

class ProposalTitleTests(SimpleTestCase):
    def test_title_ignores_logo_and_large_page_number(self):
        deck = Presentation()
        for title in ("수행 전략과 운영 계획", "투입 인력과 증빙 현황"):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index, (text, size) in enumerate((("공통회사로고",10),("03",40),(title,28))):
                shape=slide.shapes.add_textbox(Inches(1),Inches(index+1),Inches(8),Inches(1))
                shape.text=text
                shape.text_frame.paragraphs[0].runs[0].font.size=Pt(size)
            self.assertEqual(_slide_title(_slide_elements(slide)),title)
        buffer=BytesIO();deck.save(buffer)
        self.assertEqual(inspect_proposal_quality(buffer.getvalue())["duplicate_titles"],[])

    def test_explicit_title_placeholder_wins_over_larger_body(self):
        deck=Presentation()
        slide=deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text="실제 제목"
        slide.placeholders[1].text="큰 본문"
        slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.size=Pt(44)
        self.assertEqual(_slide_title(_slide_elements(slide)),"실제 제목")
