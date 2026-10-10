from io import BytesIO
import json
from unittest.mock import Mock, patch

from django.test import SimpleTestCase
from pptx import Presentation
from pptx.util import Inches

from .services.proposal_coverage import review_requirement_coverage, confirmed_quote, CoverageVerdict
from .services.proposal_detail_pages import append_detail_pages, bind_existing_details, prepare_default_template_output
from .services.proposal_output_review import output_plan, written_pages, audit_output
from .services.proposal_final_review import review_final_document


def source_result():
    prs = Presentation()
    prs.slide_width,prs.slide_height = Inches(13.333),Inches(7.5)
    page = prs.slides.add_slide(prs.slide_layouts[6])
    page.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2)).text = '원본 페이지 보존'
    buffer = BytesIO();prs.save(buffer)
    return {'file_bytes':buffer.getvalue(),'source_slide_count':1,'output_slide_count':1,
            'source_page_map':{1:1},'revision_log':[],'quality_review':{}}


def brief(number=1, **values):
    return {'id':f'R{number:04d}','requirement':'학생 250명에게 주 3회 20분 교육 제공',
            'method':'학생 250명 대상 주 3회 20분의 교육을 편성하여 운영할 계획입니다.',
            'responsible':'교육 운영 담당자','schedule':'공고 기간 내 운영',
            'deliverable':'출석부와 교육 결과보고서','verification':'출석과 수업시간을 대조하겠습니다.',
            'question':'','evidence_ids':[],'channel':'body','points':None,
            'sources':['제안요청서 3쪽'],'source_slide_numbers':[1],'output_slide_numbers':[1],**values}


class DetailPageTests(SimpleTestCase):
    def test_full_korean_and_numeric_dates_match_but_a_different_day_or_duration_does_not(self):
        numeric='교육 시작일은 2026.08.31.이며 총 15주간 운영합니다.'
        korean='교육 시작일은 2026년 8월 31일이며 총 15주간 운영합니다.'
        spaced=numeric.replace('2026.08.31','2026. 08. 31')
        for source,quote in ((numeric,korean),(korean,numeric),(spaced,korean),(korean,spaced)):
            self.assertTrue(confirmed_quote(CoverageVerdict(covered=True,slide_number=1,quote=quote),source,{1:quote}))
        for quote in (korean.replace('31일','30일'),korean.replace('15주','15개월'),korean.replace('2026년','2027년')):
            self.assertFalse(confirmed_quote(CoverageVerdict(covered=True,slide_number=1,quote=quote),numeric,{1:quote}))
        self.assertFalse(confirmed_quote(CoverageVerdict(covered=True,slide_number=1,quote=numeric),
            '2026년 8월 31개 교육 자료를 총 15주간 제공합니다.',{1:numeric}))

    def test_read_only_full_review_reconnects_actual_detail_pages_and_ignores_unrelated_ids(self):
        result=append_detail_pages(source_result(),{'items':[brief()]})
        plan={'items':[brief(output_slide_numbers=[1]),brief(2,output_slide_numbers=[1])],'notes':[]}
        bind_existing_details(result['file_bytes'],plan)
        self.assertEqual(plan['items'][0]['output_slide_numbers'],[2,1])
        self.assertEqual(plan['items'][0]['detail_slide_numbers'],[2])
        self.assertEqual(plan['detail_pages']['placements'],{'R0001':2})
        self.assertEqual(plan['detail_pages']['omitted'][0]['id'],'R0002')
        self.assertEqual(plan['items'][1]['output_slide_numbers'],[1])

    def test_custom_template_with_only_one_layout_can_add_details(self):
        original=source_result();prs=Presentation(BytesIO(original['file_bytes']))
        layouts=prs.slide_layouts
        for index in range(len(layouts)-1,0,-1):
            layouts._sldLayoutIdLst.remove(layouts._sldLayoutIdLst[index])
        buffer=BytesIO();prs.save(buffer);original['file_bytes']=buffer.getvalue()
        result=append_detail_pages(original,{'items':[brief()]})
        self.assertEqual(result['output_slide_count'],2)

    def test_existing_page_preserved_and_complete_briefs_bound_to_editable_pages(self):
        original=source_result()
        before=Presentation(BytesIO(original['file_bytes'])).slides[0]._element.xml
        plan={'items':[brief(),brief(2,channel='price'),brief(3,question='인력 배정 확인 필요')],'notes':[]}
        result=append_detail_pages(original,plan)
        prs=Presentation(BytesIO(result['file_bytes']))
        self.assertEqual(prs.slides[0]._element.xml,before)
        self.assertEqual(result['source_slide_count'],1)
        self.assertEqual(result['source_page_map'],{1:1})
        self.assertEqual(plan['detail_pages']['placed_count'],2)
        self.assertNotIn('R0002',plan['detail_pages']['placements'])
        row=plan['items'][0];number=row['detail_slide_numbers'][0]
        self.assertEqual(row['output_slide_numbers'],[number,1])
        self.assertIn(row['requirement'],written_pages(result['file_bytes'])[number])
        for key in ('method','responsible','schedule','deliverable','verification'):
            self.assertIn(row[key],written_pages(result['file_bytes'])[number])
        self.assertIn('제안요청서 3쪽',prs.slides[number-1].notes_slide.notes_text_frame.text)
        self.assertEqual(result['quality_review']['overflow_items'],[])
        self.assertEqual(result['quality_review']['small_text_items'],[])

    def test_question_follows_answer_with_explicit_point_leading(self):
        result=append_detail_pages(source_result(),{'items':[brief(question='투입 인원 확인 필요')]})
        shapes=Presentation(BytesIO(result['file_bytes'])).slides[1].shapes
        answer=next(s for s in shapes if s.name=='bid3-answer-R0001')
        question=next(s for s in shapes if s.name=='bid3-question-R0001')
        self.assertGreaterEqual(question.top,answer.top+answer.height)
        for box in (answer,question):
            self.assertTrue(all(p.line_spacing.pt==15 for p in box.text_frame.paragraphs))

    def test_page_limit_keeps_original_and_reports_every_omitted_body_condition(self):
        original=source_result();plan={'items':[brief(),brief(2)]}
        self.assertIs(append_detail_pages(original,plan,page_limit=1),original)
        self.assertEqual(plan['detail_pages']['placed_count'],0)
        self.assertEqual([r['id'] for r in plan['detail_pages']['omitted']],['R0001','R0002'])
        self.assertTrue(any('미배치' in note for note in plan['notes']))

    def test_many_conditions_pack_without_clipping_and_never_exceed_cap(self):
        rows=[brief(n,requirement=f'조건 {n}: '+('수업 및 출석 현황을 확인합니다. '*10)) for n in range(1,90)]
        plan={'items':rows}
        result=append_detail_pages(source_result(),plan,page_limit=4)
        self.assertLessEqual(result['output_slide_count'],4)
        summary=plan['detail_pages']
        self.assertEqual(summary['placed_count']+len(summary['omitted']),len(rows))
        self.assertGreater(len(summary['omitted']),0)
        self.assertEqual(result['quality_review']['overflow_items'],[])
        self.assertEqual(result['quality_review']['severe_overflow_items'],[])

    def test_oversized_condition_is_not_truncated_or_set_in_tiny_font(self):
        row=brief(requirement='긴 원문 조건 '*1000);plan={'items':[row]}
        result=append_detail_pages(source_result(),plan)
        self.assertEqual(result['output_slide_count'],1)
        self.assertEqual(plan['detail_pages']['omitted'][0]['id'],row['id'])
        self.assertEqual(plan['items'][0]['requirement'],row['requirement'])

    def test_explicit_scores_prioritize_limited_space_and_no_body_rows_add_nothing(self):
        plan={'items':[brief(1),brief(2,points=40)]}
        result=append_detail_pages(source_result(),plan,page_limit=2)
        page=written_pages(result['file_bytes'])[2]
        self.assertLess(page.index('R0002'),page.index('R0001'))
        empty={'items':[brief(channel='eligibility')]};original=source_result()
        self.assertIs(append_detail_pages(original,empty),original)

    def test_original_condition_is_visible_but_never_used_as_answer_evidence(self):
        row=brief(method='교육 일정과 교사를 배정할 계획입니다.')
        plan={'items':[row]};result=append_detail_pages(source_result(),plan)
        number=row['detail_slide_numbers'][0]
        self.assertIn(row['requirement'],written_pages(result['file_bytes'])[number])
        exported=output_plan(result['file_bytes'])
        self.assertNotIn(row['requirement'],exported['slide_changes'][number-1]['text_changes'][0]['revised_text'])
        coverage={'checks':[{'id':row['id'],'requirement':row['requirement'],
            'covered':True,'slide_number':number,'quote':row['requirement']}]}
        audited=audit_output(result['file_bytes'],{'requirement_register':{'requirements':[row]},'requirement_coverage':coverage})
        self.assertEqual(audited['matched_count'],0)

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_management_ids_and_numbered_titles_do_not_create_conflict_pairs(self,chain):
        def make(prompt,model,schema):
            def invoke(inputs):
                self.assertIn('blocks',inputs)
                return schema.model_validate({block['id']:{'kind':'other','reason':'조건 또는 설명'}
                                              for block in json.loads(inputs['blocks'])})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=make
        plan={'items':[brief(n,requirement=f'공고 조건 {n}명 교육', method='교육을 편성하여 운영합니다.',
            evidence_ids=[f'K{n}']) for n in range(1,10)]}
        result=append_detail_pages(source_result(),plan)
        report=review_final_document(result['file_bytes'],'',Mock())
        self.assertEqual(report['conflicts'],[])
        self.assertEqual(report['failures'],[])
        self.assertGreater(report['reviewed_block_count'],len(plan['items']))

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_field_labels_alone_do_not_connect_unrelated_numeric_subjects(self,chain):
        captured=[]
        def make(prompt,model,schema):
            def invoke(inputs):
                if 'candidate' in inputs:
                    return schema.model_validate({'verdict':'conflict','same_subject':True,'same_metric':True,
                        'same_time_scope':True,'reason':json.loads(inputs['candidate'])['reason']})
                if 'blocks' in inputs:
                    return schema.model_validate({block['id']:{'kind':'future_plan','reason':'계획'}
                                                  for block in json.loads(inputs['blocks'])})
                captured.extend(json.loads(inputs['pairs']))
                return schema.model_validate({'items':[]})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=make
        prs=Presentation()
        page=prs.slides.add_slide(prs.slide_layouts[6])
        for i,text in enumerate(('수행안: 학생 250명 교육을 편성합니다.',
                                '수행안: 라이선스 100개 구입을 제안합니다.',
                                '검증: 학생 300명 교육을 편성합니다.')):
            page.shapes.add_textbox(0,Inches(i),Inches(8),Inches(1)).text=text
        buffer=BytesIO();prs.save(buffer)
        review_final_document(buffer.getvalue(),'',Mock())
        self.assertEqual(len(captured),1)
        self.assertIn('250명',captured[0]['first_quote'])
        self.assertIn('300명',captured[0]['second_quote'])

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_invalid_conflict_batch_rechecks_each_pair_once_and_keeps_failed_pairs_visible(self,chain):
        retried=[]
        def make(prompt,model,schema):
            def invoke(inputs):
                if 'candidate' in inputs:
                    return schema.model_validate({'verdict':'conflict','same_subject':True,'same_metric':True,
                        'same_time_scope':True,'reason':json.loads(inputs['candidate'])['reason']})
                if 'blocks' in inputs:
                    return schema.model_validate({block['id']:{'kind':'other','reason':'설명'} for block in json.loads(inputs['blocks'])})
                pairs=json.loads(inputs['pairs'])
                if len(pairs)>1:
                    raise ValueError('Invalid structured reply')
                retried.extend(pairs)
                if '400명' in pairs[0]['second_quote']:
                    raise ValueError('Still invalid')
                return schema.model_validate({'items':[{**pairs[0],'reason':'동일 교육 인원 불일치'}]})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=make
        prs=Presentation()
        for text in ('총 교육 인원 250명','총 교육 인원 300명','총 교육 인원 400명'):
            prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0,0,Inches(8),Inches(2)).text=text
        buffer=BytesIO();prs.save(buffer)
        report=review_final_document(buffer.getvalue(),'',Mock())
        self.assertEqual(len(retried),3)
        self.assertEqual(report['conflict_retry_pair_count'],3)
        self.assertEqual(len(report['conflicts']),1)
        self.assertEqual(len(report['failures']),2)

    @patch('bids.services.proposal_coverage.structured_chain')
    def test_reviewer_receives_detail_destination_before_unrelated_relevant_pages(self,chain):
        seen=[]
        def make(prompt,model,schema):
            def invoke(inputs):
                seen.extend(json.loads(inputs['review_context']))
                return schema.model_validate({key:{'covered':False,'reason':'검토 필요'} for key in schema.model_fields})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=make
        plan={'slide_changes':[{'slide_number':n,'action':'UPDATE','text_changes':[{'revised_text':'학생 교육 운영'}]} for n in range(1,6)],
              'writing_plan':{'items':[{'id':'R0001','output_slide_numbers':[5,3]}]}}
        review_requirement_coverage({'requirements':[brief()]},plan,Mock())
        self.assertEqual([r['slide_number'] for r in seen[0]['candidates']],[5,3])

    @patch('bids.services.proposal_coverage.structured_chain')
    def test_copied_excerpt_wrapper_is_removed_only_when_remaining_actual_quote_passes(self,chain):
        def make(prompt,model,schema):
            return Mock(invoke=Mock(return_value=schema.model_validate({
                'R0001':{'covered':True,'quote':'[21f902141505] 용역기간은 착수일로부터 90일간입니다.'},
                'R0002':{'covered':True,'quote':'[21f902141505] 존재하지 않는 90일 문구입니다.'},
                'R0003':{'covered':True,'quote':'[21f902141505] 용역기간은 착수일로부터 90일간입니다.'}})))
        chain.side_effect=make
        prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0,0,Inches(8),Inches(2)).text='용역기간은 착수일로부터 90일간입니다.'
        buffer=BytesIO();prs.save(buffer)
        report=review_requirement_coverage({'requirements':[
            {'requirement':'착수일로부터 90일'}, {'requirement':'착수일로부터 90일'},
            {'requirement':'착수일로부터 90개월'}]},output_plan(buffer.getvalue()),Mock())
        self.assertEqual([r['covered'] for r in report['checks']],[True,False,False])
        self.assertEqual(report['checks'][0]['quote'],'용역기간은 착수일로부터 90일간입니다.')

    @patch('bids.services.proposal_coverage.structured_chain')
    def test_native_answer_is_sent_whole_without_other_cards_or_source_echo(self,chain):
        seen=[]
        def make(prompt,model,schema):
            def invoke(inputs):
                seen.extend(json.loads(inputs['review_context']))
                return schema.model_validate({'R0001':{'covered':False,'reason':'별도 검토'}})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=make
        row=brief(method='학습자 교육 운영 방법 '*24,verification='끝부분 검증 방법을 확인합니다.',question='투입 가능 인원을 확인해야 합니다.')
        result=append_detail_pages(source_result(),{'items':[row,brief(2,method='별도 조건의 답변')]})
        plan=output_plan(result['file_bytes']);plan['writing_plan']={'items':[row]}
        review_requirement_coverage({'requirements':[row]},plan,Mock())
        text=seen[0]['candidates'][0]['text']
        self.assertIn(row['method'],text)
        self.assertIn(row['verification'],text)
        self.assertIn(row['question'],text)
        self.assertNotIn('별도 조건의 답변',text)
        self.assertNotIn('공고 조건',text)

    def test_fresh_builtin_schedule_removes_invented_bars_and_keeps_source_dates(self):
        result=source_result();prs=Presentation(BytesIO(result['file_bytes']));page=prs.slides[0]
        for name,text in [('schedule-header',''),('week-0','1주'),('schedule-bar-0',''),
                          ('phase-output-0','주요 산출물\n[운영 설계서]')]:
            shape=page.shapes.add_textbox(0,0,Inches(2),Inches(1));shape.name=name;shape.text=text
        buffer=BytesIO();prs.save(buffer);result['file_bytes']=buffer.getvalue()
        original=result['file_bytes']
        clean=prepare_default_template_output(result,{'requirements':[
            {'requirement':'교육기간: 2026.08.31.~12.13. 총 15주','sources':['공고 2쪽']},
            {'requirement':'개인정보 누출 시 사업 기간 또는 종료 후라도 제안사가 책임을 짐'}]})
        page=Presentation(BytesIO(clean['file_bytes'])).slides[0]
        self.assertFalse(any(s.name.startswith(('week-','schedule-bar-')) for s in page.shapes))
        self.assertIn('2026.08.31.~12.13.', '\n'.join(s.text for s in page.shapes if s.has_text_frame))
        self.assertNotIn('개인정보 누출','\n'.join(s.text for s in page.shapes if s.has_text_frame))
        self.assertIn('공고 2쪽',page.notes_slide.notes_text_frame.text)
        self.assertEqual(next(s.text for s in page.shapes if s.name=='phase-output-0').strip(),'주요 산출물\n운영 설계서')
        self.assertEqual(result['file_bytes'],original)
        self.assertEqual(clean['quality_review']['timeline_range_items'],[])

    def test_unrecognized_schedule_is_preserved(self):
        result=source_result()
        clean=prepare_default_template_output(result,{'requirements':[]})
        self.assertEqual(written_pages(clean['file_bytes']),written_pages(result['file_bytes']))
