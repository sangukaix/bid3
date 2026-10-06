from unittest.mock import patch
from django.test import SimpleTestCase
from .services.text_geometry import fitting_size, measure_text
from .services.local_context import structured_chain
from .services.llm import OllamaChatModel
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnableLambda
from pydantic import BaseModel
from .services.context_tables import packed_json, unpack, TABLE, MAPPING
import json


class GeometryTests(SimpleTestCase):
    def test_long_korean_text_does_not_fit_but_readable_short_text_fits(self):
        element={'width':220,'height':55,'font_size':20,'font_family':'Malgun Gothic'}
        self.assertTrue(measure_text('실습 결과와 회고',element)['fits'])
        self.assertIsNone(fitting_size('구현과 검증 내용을 자세히 설명합니다. '*100,element)[0])

    def test_wrapping_and_small_labels_are_preserved(self):
        element={'width':150,'height':70,'font_size':10,'wrap':True}
        size,result=fitting_size('출처: 수업 실습 기록',element)
        self.assertEqual(size,10)
        self.assertTrue(result['estimated'])
        self.assertGreater(measure_text('첫 줄\n둘째 줄',element)['line_count'],1)

    def test_disabled_wrap_detects_horizontal_overflow(self):
        element={'width':40,'height':200,'font_size':20,'wrap':False}
        self.assertFalse(measure_text('제목이 너무 깁니다',element)['fits'])

    def test_fixed_line_spacing_does_not_shrink_with_font(self):
        element={'width':300,'height':55,'font_size':18,'line_height_pt':30,'wrap':False}
        self.assertFalse(measure_text('첫 줄\n둘째 줄',element,11)['fits'])

    def test_merged_cell_uses_combined_width_and_paragraph_font(self):
        from io import BytesIO
        from pptx import Presentation
        from pptx.util import Inches,Pt
        from .services.proposal_pptx_renderer import _slide_elements
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        table=slide.shapes.add_table(1,2,Inches(1),Inches(1),Inches(4),Inches(1)).table
        table.cell(0,0).merge(table.cell(0,1));cell=table.cell(0,0)
        cell.text='합쳐진 셀';cell.text_frame.paragraphs[0].font.size=Pt(22)
        buffer=BytesIO();prs.save(buffer)
        elements=_slide_elements(Presentation(BytesIO(buffer.getvalue())).slides[0])
        self.assertEqual(len(elements),1);self.assertEqual(elements[0]['width'],288)
        self.assertEqual(elements[0]['font_size'],22)


class PackedContextTests(SimpleTestCase):
    def test_roundtrip_keeps_numbers_types_reserved_keys_and_source_strings(self):
        rows=[{'long_source_id':f'자료:{i}','numeric_value':123456789.01,'nullable':None,
               'boolean':False,'string_zero':'0','reserved':{TABLE:'자료',MAPPING:'키'}} for i in range(20)]
        wire=packed_json(rows)
        self.assertEqual(unpack(json.loads(wire)),rows)
        self.assertLess(len(wire.encode()),len(json.dumps(rows,ensure_ascii=False).encode()))
        self.assertIs(unpack(json.loads(wire))[0]['boolean'],False)
        self.assertIsInstance(unpack(json.loads(wire))[0]['string_zero'],str)

    def test_heterogeneous_rows_do_not_fill_missing_fields(self):
        rows=[{}, {'a':None}, {'b':0}, {'a':False}]
        self.assertEqual(unpack(json.loads(packed_json(rows))),rows)

    def test_local_proposal_receives_lossless_inventory_and_reading_rule(self):
        class Output(BaseModel):
            okay: bool
        rows=[{'long_source_id':f'자료:{i}','full_original_text':'대출은 미구현; 성능 미측정',
               'original_numeric_value':90,'not_a_measured_result':False} for i in range(30)]
        inputs={'slide_inventory':json.dumps(rows,ensure_ascii=False),'instruction':'사용자 지침을 그대로 유지'}
        before=dict(inputs);captured=[]
        def invoke(value):
            captured.extend(value.to_messages() if hasattr(value,'to_messages') else value)
            return Output(okay=True)
        prompt=ChatPromptTemplate.from_messages([('system','한글로 작성'),('human','{instruction}\n{slide_inventory}')])
        with patch.object(OllamaChatModel,'with_structured_output',return_value=RunnableLambda(invoke)):
            result=structured_chain(prompt,OllamaChatModel(model='gemma4:26b',num_ctx=32768,num_predict=800),Output).invoke(inputs)
        self.assertTrue(result.okay);self.assertEqual(inputs,before)
        self.assertIn('__bid3_table_v1__',captured[0].content)
        human=captured[-1].content
        self.assertTrue(human.startswith(inputs['instruction']+'\n'))
        self.assertEqual(unpack(json.loads(human.split('\n',1)[1])),rows)
