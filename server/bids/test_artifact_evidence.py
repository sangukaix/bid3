from io import BytesIO
from unittest.mock import patch
from django.test import SimpleTestCase
from pptx import Presentation
from pptx.util import Inches
from .services.proposal_coverage import CoverageVerdict, confirmed_quote, review_requirement_coverage
from .services.proposal_output_review import audit_output, output_plan, refresh_output_review, written_pages


def deck(*texts):
    prs=Presentation()
    for text in texts:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0,0,Inches(8),Inches(3)).text=text
    buffer=BytesIO();prs.save(buffer)
    return buffer.getvalue()


class ArtifactEvidenceTests(SimpleTestCase):
    def plan(self,quote,requirement='250명에게 30분 수업',page=1):
        return {'requirement_register':{'requirements':[{'requirement':requirement,'priority':'필수',
            'evaluation_points':'10점','sources':['요청서 4쪽']}]},
            'requirement_coverage':{'covered_count':1,'checks':[{'id':'R0001','requirement':requirement,
                'covered':True,'slide_number':page,'quote':quote}]}}

    def test_numbers_with_different_units_dates_or_uncertainty_do_not_pass(self):
        for requirement,quote in [('250명 30분 수업','250명 30시간 수업을 운영합니다.'),
            ('2027.01.02 종료','2027.02.01 종료하며 결과를 보고합니다.'),
            ('250명 30분 수업','250명 30분 수업 운영은 확인 필요합니다.')]:
            with self.subTest(requirement=requirement):
                self.assertFalse(confirmed_quote(CoverageVerdict(covered=True,slide_number=1,quote=quote),requirement,{1:quote}))
        quote='2027.1.2에 종료하며 0.50시간 교육을 제공합니다.'
        self.assertTrue(confirmed_quote(CoverageVerdict(covered=True,slide_number=1,quote=quote),
            '2027.01.02 종료, 0.5시간 교육', {1:quote}))

    def test_output_remaps_a_quote_after_page_insertion_and_preserves_rfp_sources(self):
        quote='250명에게 30분 수업을 운영합니다.'
        review=audit_output(deck('새 표지',quote),self.plan(quote))
        self.assertEqual(review['matched_count'],1)
        self.assertEqual(review['checks'][0]['slide_number'],2)
        self.assertEqual(review['checks'][0]['evaluation_points'],'10점')
        self.assertEqual(review['checks'][0]['sources'],['요청서 4쪽'])

    def test_changed_or_rejected_text_invalidates_both_match_and_old_counter(self):
        plan=self.plan('250명에게 30분 수업을 운영합니다.')
        report=refresh_output_review(deck('실제 적용된 기존 문구만 남아 있습니다.'),plan)
        self.assertEqual(report['matched_count'],0)
        self.assertEqual(plan['requirement_coverage']['covered_count'],0)
        self.assertEqual(plan['requirement_coverage']['total_count'],1)
        self.assertEqual(plan['requirement_coverage']['checks'][0]['quote'],'')

    def test_removed_page_or_ambiguous_renumbering_does_not_invent_a_location(self):
        quote='250명에게 30분 수업을 운영합니다.'
        review=audit_output(deck(quote,quote),self.plan(quote,page=99))
        self.assertEqual(review['matched_count'],0)
        self.assertIsNone(review['checks'][0]['slide_number'])

    def test_reused_id_with_a_different_requirement_does_not_reuse_old_verdict(self):
        quote='250명에게 30분 수업을 운영합니다.'
        plan=self.plan(quote)
        plan['requirement_register']['requirements'][0]['requirement']='250명에게 30분 사후 상담'
        self.assertEqual(audit_output(deck(quote),plan)['matched_count'],0)

    def test_legacy_aggregate_without_register_is_explicitly_unreviewed(self):
        report=audit_output(deck('회사 실적 자료 확인 필요'),{'requirement_coverage':{'covered_count':99}})
        self.assertFalse(report['source_register_available'])
        self.assertTrue(report['review_notes'])
        self.assertEqual(report['open_text_items'][0]['slide_number'],1)

    def test_tables_groups_and_added_pages_are_read_but_notes_are_not_evidence(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        group=slide.shapes.add_group_shape()
        group.shapes.add_textbox(0,0,Inches(3),Inches(1)).text='그룹 안의 실제 문장'
        table=slide.shapes.add_table(1,2,0,Inches(2),Inches(6),Inches(1)).table
        table.cell(0,0).merge(table.cell(0,1));table.cell(0,0).text='표 안의 실제 조건'
        slide.notes_slide.notes_text_frame.text='검수 계획에만 있는 250명 30분'
        buffer=BytesIO();prs.save(buffer)
        pages=written_pages(buffer.getvalue())
        self.assertIn('그룹 안의 실제 문장',pages[1]);self.assertIn('표 안의 실제 조건',pages[1])
        self.assertEqual(pages[1].count('표 안의 실제 조건'),1)
        self.assertNotIn('250명',pages[1])
        self.assertEqual(output_plan(deck('첫 장','추가된 장'))['slide_changes'][1]['slide_number'],2)

    @patch('bids.services.proposal_coverage.structured_chain')
    def test_semantic_reviewer_receives_actual_output_not_unapplied_plan(self,factory):
        from unittest.mock import Mock
        captured=[]
        def build(prompt,model,schema):
            def invoke(values):
                import json
                payload=json.loads(values['review_context']);captured.extend(payload)
                return schema.model_validate({item['id']:{'covered':False} for item in payload})
            return Mock(invoke=invoke)
        factory.side_effect=build
        register={'requirements':[{'requirement':'추가된 운영 방법'}]}
        review_requirement_coverage(register,output_plan(deck('표지','추가된 운영 방법을 설명합니다.')),None)
        self.assertTrue(any('추가된 운영 방법을 설명합니다.' in item['text'] for item in captured[0]['candidates']))
