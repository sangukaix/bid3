from copy import deepcopy
from contextlib import ExitStack
from hashlib import sha256
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import Mock, patch
from django.test import SimpleTestCase
from django.test import TestCase, override_settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.utils import timezone
from rest_framework.test import APIClient
from tempfile import TemporaryDirectory
from .models import BidNotice, SavedBid, BidProposal, ProposalTask
from .services.proposal_tasks import enqueue, run_task, recover_dead_task, process_alive
from pptx import Presentation
from pptx.util import Inches
from .services.proposal_planning import channel_for, explicit_points, build_writing_plan, bind_output_pages
from .services.proposal_final_review import review_final_document, invalidate_final_review, repair_once, digest, confirm_conflicts, comparable_quantities
from .services.proposal_evidence import evidence_units


def deck(*texts):
    presentation=Presentation()
    for text in texts:
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0,0,Inches(8),Inches(4)).text=text
    buffer=BytesIO();presentation.save(buffer);return buffer.getvalue()


class WritingPlanTests(SimpleTestCase):
    @patch('bids.services.proposal_planning.structured_chain')
    def test_invalid_brief_format_retries_once_without_dropping_the_condition(self,chain):
        schemas=[]
        def factory(prompt,model,schema):
            schemas.append(schema)
            if len(schemas)==1: return Mock(invoke=Mock(side_effect=ValueError('Too long')))
            return Mock(invoke=Mock(return_value=schema.model_validate({'R0001':{
                'method':'교육을 편성합니다.','responsible':'교육 운영팀','schedule':'교육 기간 내',
                'deliverable':'출석부','verification':'출석 대조','question':'','evidence_ids':[]}})))
        chain.side_effect=factory
        result=build_writing_plan({'requirements':[{'requirement':'교육 운영'}]},[],'',Mock())
        self.assertEqual(len(schemas),2)
        self.assertEqual(result['items'][0]['method'],'교육을 편성합니다.')
        self.assertNotIn('작성 설계 실패',result['items'][0]['question'])

    def test_submission_channels_do_not_claim_narrative_completes_proof_or_price(self):
        for text,expected in [('입찰 참가 자격 실적증명서 제출','eligibility'),('가격제안 별도 제출','price'),
                              ('사업수행계획','body'),('제출 방법 방문접수','manual'),
                              ('입찰 자격 및 가격입찰','manual')]:
            self.assertEqual(channel_for({'requirement':text}),expected)
        self.assertEqual(channel_for({'requirement':'인력 현황','form_name':'별지 1호'}),'form')

    def test_scores_do_not_turn_ranges_multi_scores_or_digits_into_weights(self):
        self.assertEqual(explicit_points('배점 12.5점'),12.5)
        for value in ('10~20점','2점/3점','가격 300만원','10-20점','합계 100점'):
            self.assertIsNone(explicit_points(value))

    @patch('bids.services.proposal_planning.structured_chain')
    def test_every_condition_kept_bad_ids_removed_and_high_score_gets_more_detail(self,chain):
        rows={'requirements':[{'requirement':'운영 방안','evaluation_points':'40점'},
                             {'requirement':'보고서','evaluation_points':'5점'},
                             {'requirement':'가격입찰'}]}
        inventory=[{'slide_number':n,'role':'body','elements':[{}],'title':'운영 방안'} for n in range(2,8)]
        def make(prompt,model,schema):
            values={key:{'method':'제안 수행안','responsible':'역할 배정 필요','schedule':'원문 기간',
                'deliverable':'보고서','verification':'담당자 검토','evidence_ids':['K1','K999'],'question':''}
                for key in schema.model_fields}
            return Mock(invoke=Mock(return_value=schema.model_validate(values)))
        chain.side_effect=make
        result=build_writing_plan(rows,inventory,'[검토 완료 근거 K1 | 실적]\n확인된 내용',Mock())
        high,low,price=result['items']
        self.assertEqual(set(chain.call_args.args[2].model_fields),{'R0001','R0002'})
        self.assertEqual(len(result['items']),3)
        self.assertGreater(len(high['source_slide_numbers']),len(low['source_slide_numbers']))
        self.assertEqual(high['evidence_ids'],['K1']);self.assertEqual(price['source_slide_numbers'],[])
        bind_output_pages(result,{2:3,3:None,4:5,5:6,6:7,7:8})
        self.assertNotIn(None,high['output_slide_numbers'])

    @patch('bids.services.proposal_planning.structured_chain',side_effect=TimeoutError())
    def test_failed_batch_and_no_content_pages_stay_explicit(self,chain):
        result=build_writing_plan({'requirements':[{'requirement':'운영 방안'}]},[], '',Mock())
        self.assertEqual(len(result['items']),1)
        self.assertIn('실패',result['items'][0]['question'])
        self.assertIn('페이지가 없습니다',result['items'][0]['question'])

    def test_reviewed_company_item_qualifications_are_not_fragmented(self):
        text='[검토 완료 근거 K1 | 인력]\n'+('인력 근거 '*500)+'\n확정 아님\n\n[검토 완료 근거 K2 | 실적]\n실적 근거'
        units=evidence_units(text,'company_knowledge_context')
        self.assertEqual([unit.id for unit in units],['K1','K2'])
        self.assertIn('확정 아님',units[0].text)

    def test_removed_pages_are_not_exposed_as_output_destinations(self):
        plan={'items':[{'source_slide_numbers':[1,2,3]}]}
        bind_output_pages(plan,{1:1,2:None,3:2})
        self.assertEqual(plan['items'][0]['output_slide_numbers'],[1,2])

    @patch('bids.services.proposal_planning.structured_chain')
    def test_separate_submission_tasks_preserve_conditions_without_ai_or_body_pages(self,chain):
        rows={'requirements':[{'requirement':'가격입찰 부가세 10% 포함'},
                             {'requirement':'입찰 참가 자격 면허 증빙 제출'},
                             {'requirement':'작성 서류 제출','form_name':'별지 1호'},
                             {'requirement':'제출 방법 방문접수'}]}
        plan=build_writing_plan(rows,[], '',Mock())
        chain.assert_not_called()
        self.assertEqual([r['requirement'] for r in plan['items']],[r['requirement'] for r in rows['requirements']])
        self.assertEqual([r['channel'] for r in plan['items']],['price','eligibility','form','manual'])
        self.assertEqual(plan['items'][2]['deliverable'],'별지 1호')
        self.assertTrue(all(r['question'] and not r['evidence_ids'] and not r['source_slide_numbers'] for r in plan['items']))


class FinalDocumentTests(SimpleTestCase):
    @patch('bids.services.proposal_final_review.structured_chain')
    def test_batch_never_has_more_candidates_than_the_response_schema_can_report(self,chain):
        sizes=[]
        def factory(prompt,model,schema):
            def invoke(inputs):
                if 'blocks' in inputs:
                    return schema.model_validate({b['id']:{'kind':'other','reason':'설명'} for b in json.loads(inputs['blocks'])})
                if 'candidate' in inputs:
                    return schema.model_validate({'verdict':'conflict','same_subject':True,'same_metric':True,
                        'same_time_scope':True,'reason':'동일 교육 인원 불일치'})
                pairs=json.loads(inputs['pairs']);sizes.append(len(pairs))
                # A constrained model cannot emit more than the schema permits.
                return schema.model_validate({'items':[{**pair,'reason':'동일 교육 인원 불일치'} for pair in pairs[:5]]})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=factory
        report=review_final_document(deck(*(f'총 교육 인원 {n}명' for n in range(250,257))),'',Mock())
        self.assertEqual(len(report['conflicts']),21)
        self.assertLessEqual(max(sizes),5)
        self.assertEqual(report['conflict_retry_pair_count'],0)

    def test_unit_families_preserve_wrong_duration_units_but_do_not_equate_people_and_classes(self):
        self.assertTrue(comparable_quantities('총 교육 시간 30분','총 교육 시간 30시간'))
        self.assertTrue(comparable_quantities('용역 기간 90일','용역 기간 90개월'))
        self.assertFalse(comparable_quantities('동시접속 150명','동시접속 40class'))
        self.assertFalse(comparable_quantities('사업 예산 250만원','교육 인원 250명'))

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_generic_schedule_language_does_not_relate_unidentified_tasks(self,chain):
        def factory(prompt,model,schema):
            def invoke(inputs):
                self.assertIn('blocks',inputs)
                return schema.model_validate({b['id']:{'kind':'future_plan','reason':'계획'} for b in json.loads(inputs['blocks'])})
            return Mock(invoke=Mock(side_effect=invoke))
        chain.side_effect=factory
        report=review_final_document(deck('일정: 사업 착수 후 7일 이내 구축 완료',
            '일정: 착수일로부터 교육 시작일(2026.08.31) 전까지 완료'),'',Mock())
        self.assertEqual(report['conflicts'],[])
        self.assertEqual(report['failures'],[])

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_conflict_requires_same_subject_metric_and_scope_even_if_model_label_disagrees(self,chain):
        from .services.proposal_final_review import ConflictDecision
        chain.return_value.invoke.side_effect=[ValueError('Invalid reply'),ConflictDecision(
            verdict='conflict',same_subject=False,same_metric=True,same_time_scope=True,reason='서로 다른 과업입니다.')]
        candidate={'first_page':1,'first_quote':'시스템 구축 7일','second_page':2,'second_quote':'강사 구성 14일','reason':'후보'}
        accepted,dismissed,failures=confirm_conflicts([candidate],{1:'시스템 구축 7일',2:'강사 구성 14일'}, {},Mock())
        self.assertEqual(accepted,[])
        self.assertEqual(len(dismissed),1)
        self.assertEqual(failures,[])
        self.assertEqual(chain.call_count,2)

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_conflict_recheck_preserves_uncertainty_failures_and_dismissal_reason(self,chain):
        values=[{'verdict':'compatible','same_subject':True,'same_metric':True,'same_time_scope':False,
                 'reason':'전체 사업 기간과 교육 기간은 다른 범위입니다.'},
                {'verdict':'uncertain','same_subject':None,'same_metric':True,'same_time_scope':None,
                 'reason':'동일 시스템인지 원문으로 확정되지 않습니다.'},TimeoutError()]
        def make(prompt,model,schema):
            value=values.pop(0)
            if isinstance(value,Exception): return Mock(invoke=Mock(side_effect=value))
            return Mock(invoke=Mock(return_value=schema.model_validate(value)))
        chain.side_effect=make
        candidates=[{'first_page':1,'first_quote':'전체 사업 20주','second_page':2,
                     'second_quote':'교육 운영 15주','reason':f'후보 {n}'} for n in range(3)]
        accepted,dismissed,failures=confirm_conflicts(candidates,{1:'전체 사업 20주',2:'교육 운영 15주'}, {},Mock())
        self.assertEqual(len(dismissed),1)
        self.assertIn('다른 범위',dismissed[0]['review_reason'])
        self.assertEqual(accepted[0]['verification'],'uncertain')
        self.assertEqual(accepted[1],candidates[2])
        self.assertEqual(len(failures),1)

    def chain(self,kind='company_fact',supported=False,quote='',conflicts=None):
        def make(prompt,model,schema):
            def invoke(values):
                if 'candidate' in values:
                    return schema.model_validate({'verdict':'conflict','same_subject':True,'same_metric':True,
                        'same_time_scope':True,'reason':json.loads(values['candidate'])['reason']})
                if 'blocks' in values:
                    return schema.model_validate({b['id']:{'kind':kind,'supported':supported,
                        'evidence_quote':quote,'reason':'회사 근거 대조'} for b in json.loads(values['blocks'])})
                return schema.model_validate({'items':conflicts or []})
            return Mock(invoke=Mock(side_effect=invoke))
        return make

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_unchanged_added_and_non_regex_company_facts_are_all_inspected(self,chain):
        chain.side_effect=self.chain()
        report=review_final_document(deck('기존 고객사는 전국 10곳입니다.','수정 페이지','추가 페이지'),'',Mock())
        self.assertEqual(report['actual_slide_count'],3)
        self.assertEqual(report['reviewed_block_count'],3)
        self.assertEqual({item['slide_number'] for item in report['company_claim_review']['items']},{1,2,3})
        self.assertEqual(report['review_required_count'],3)

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_only_exact_numeric_unit_matching_approved_evidence_is_connected(self,chain):
        claim='교육 250명에게 30시간을 제공한 실적이 있습니다.'
        evidence='[검토 완료 근거 K1 | 실적]\n교육 250명에게 30분을 제공한 실적이 있습니다.'
        chain.side_effect=self.chain(supported=True,quote=claim)
        report=review_final_document(deck(claim),evidence,Mock())
        self.assertEqual(report['review_required_count'],1)
        chain.side_effect=self.chain(supported=True,quote='교육 250명에게 30분을 제공한 실적이 있습니다.')
        report=review_final_document(deck(claim),evidence,Mock())
        self.assertEqual(report['review_required_count'],1)
        report=review_final_document(deck('교육 250명에게 30분을 제공한 실적이 있습니다.'),evidence,Mock())
        self.assertEqual(report['review_required_count'],0)

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_future_plan_not_misreported_as_company_proof(self,chain):
        chain.side_effect=self.chain(kind='future_plan')
        report=review_final_document(deck('교육 인력 3명을 배정할 예정입니다.'),'',Mock())
        self.assertEqual(report['company_claim_review']['items'],[])

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_evidence_metadata_alone_cannot_prove_company_facts(self,chain):
        claim='2026년 인증을 보유하고 있습니다.'
        chain.side_effect=self.chain(supported=True,quote='2026년 인증')
        report=review_final_document(deck(claim),'[검토 완료 근거 K1 | 2026년 인증]\n원문 확인 필요',Mock())
        self.assertEqual(report['review_required_count'],1)

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_equal_digits_with_different_units_are_still_conflict_candidates(self,chain):
        conflict={'first_page':1,'first_quote':'총 교육 시간 30분','second_page':2,
                  'second_quote':'총 교육 시간 30시간','reason':'같은 교육 시간의 단위 불일치'}
        chain.side_effect=self.chain(kind='other',conflicts=[conflict])
        report=review_final_document(deck(conflict['first_quote'],conflict['second_quote']),'',Mock())
        self.assertEqual(report['conflicts'],[conflict])

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_rfp_conditions_are_classification_context_but_never_company_proof(self,chain):
        claim='당사는 교육 250명을 수행한 실적이 있습니다.'
        register={'requirements':[{'requirement':'교육 250명 수행'}]}
        calls=[]
        def factory(prompt,model,schema):
            proxy=self.chain(supported=True,quote='교육 250명 수행')(prompt,model,schema)
            invoke=proxy.invoke.side_effect
            proxy.invoke.side_effect=lambda values:(calls.append(values) or invoke(values))
            return proxy
        chain.side_effect=factory
        content=deck(claim)
        report=review_final_document(content,'',Mock(),register)
        self.assertIn('교육 250명 수행',calls[0]['requirements'])
        self.assertNotIn('교육 250명 수행',calls[0]['company_knowledge_context'])
        self.assertEqual(report['review_required_count'],1)
        plan={'requirement_register':register,'final_document_review':report}
        invalidate_final_review(content,plan);self.assertFalse(report['stale'])
        plan['requirement_register']={'requirements':[{'requirement':'교육 300명 수행'}]}
        invalidate_final_review(content,plan);self.assertTrue(report['stale'])

    @patch('bids.services.proposal_final_review.structured_chain',side_effect=TimeoutError())
    def test_failed_classification_is_never_all_clear(self,chain):
        report=review_final_document(deck('본문입니다.'),'',Mock())
        self.assertEqual(len(report['failures']),1)
        self.assertGreater(report['review_required_count'],0)

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_conflicting_numbers_require_actual_quote_and_page_pairs(self,chain):
        conflict={'first_page':1,'first_quote':'총 교육 인원 250명','second_page':2,'second_quote':'총 교육 인원 300명','reason':'동일 총원 불일치'}
        chain.side_effect=self.chain(kind='other',conflicts=[conflict])
        report=review_final_document(deck('총 교육 인원 250명','총 교육 인원 300명'),'',Mock())
        self.assertEqual(report['conflicts'],[conflict])
        conflict['second_quote']='총 교육 인원 500명'
        report=review_final_document(deck('총 교육 인원 250명','총 교육 인원 300명'),'',Mock())
        self.assertEqual(report['conflicts'],[])

    def test_file_or_company_evidence_change_invalidates_without_model(self):
        content=deck('운영 계획')
        plan={'final_document_review':{'file_sha256':sha256(content).hexdigest(),'company_evidence_sha256':digest('approved'),'stale':False}}
        invalidate_final_review(content,plan,'approved');self.assertFalse(plan['final_document_review']['stale'])
        invalidate_final_review(content,plan,'expired');self.assertTrue(plan['final_document_review']['stale'])
        plan['final_document_review']['stale']=False
        invalidate_final_review(deck('수정 운영 계획'),plan);self.assertTrue(plan['final_document_review']['stale'])

    @patch('bids.services.proposal_final_review.structured_chain')
    def test_fixed_group_table_long_text_and_notes_have_correct_review_scope(self,chain):
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        group=page.shapes.add_group_shape();group.shapes.add_textbox(0,0,Inches(6),Inches(2)).text='그룹 회사 주장'
        fixed=page.shapes.add_textbox(0,0,Inches(6),Inches(2));fixed.name='bid3-fixed-private';fixed.text='고정 회사 주장'
        page.shapes.add_table(1,1,0,0,Inches(6),Inches(2)).table.cell(0,0).text='표 안 회사 주장'
        page.shapes.add_textbox(0,0,Inches(6),Inches(2)).text='가'*3000+'끝의 회사 주장'
        page.notes_slide.notes_text_frame.text='노트는 제출 본문 아님'
        buffer=BytesIO();prs.save(buffer);chain.side_effect=self.chain()
        report=review_final_document(buffer.getvalue(),'',Mock())
        claims=[item['claim'] for item in report['company_claim_review']['items']]
        self.assertIn('그룹 회사 주장',claims);self.assertIn('고정 회사 주장',claims);self.assertIn('표 안 회사 주장',claims)
        self.assertTrue(any('끝의 회사 주장' in text for text in claims))
        self.assertFalse(any('노트는' in text for text in claims))


class BoundedRepairTests(SimpleTestCase):
    def fixture(self):
        content=deck(*['미확인 회사 사실' for _ in range(5)])
        quality={'severe_overflow_items':[],'unresolved_placeholders':[],'template_leftovers':[]}
        file={'file_bytes':content,'source_slide_count':5,'revision_log':[], 'quality_review':quality}
        report={'company_claim_review':{'items':[{'slide_number':n,'claim':'사실','reason':'자료 없음','status':'review_required'} for n in range(1,6)]},'conflicts':[],'failures':[],'review_required_count':5}
        plan={'final_document_review':report,'requirement_coverage':{'checks':[]},
              'output_review':{'review_required_count':0,'open_text_items':[],'matched_count':0}}
        return file,plan

    @patch('bids.services.proposal_final_review.verify_artifact')
    @patch('bids.services.proposal_final_review.build_proposal_pptx')
    def test_attempt_is_once_at_most_three_pages_and_unrelated_targets_ignored(self,render,verify):
        file,plan=self.fixture();candidate=deepcopy(file);candidate['file_bytes']=deck('개선안');render.return_value=candidate
        def verified(content,p,*args):
            p['final_document_review']['review_required_count']=2
        verify.side_effect=verified
        def edit(slide,inputs):
            return SimpleNamespace(model_dump=lambda:{'slide_changes':[{'slide_number':slide['slide_number'],'action':'UPDATE','text_changes':[
                {'target':slide['elements'][0]['target'],'revised_text':'확인 필요'}, {'target':'forged','revised_text':'삭제'}]},
                {'slide_number':99,'action':'REMOVE','text_changes':[]}],'added_slides':[{'title':'금지'}]})
        output=repair_once(file,plan,Mock(),'','',Mock(),Mock(),edit)
        self.assertIs(output,candidate);self.assertTrue(plan['bounded_repair']['accepted'])
        self.assertEqual(plan['bounded_repair']['pages'],[1,2,3]);verify.assert_called_once()
        changes=render.call_args.args[2]['slide_changes']
        self.assertEqual(len(changes),3)
        self.assertTrue(all(len(change['text_changes'])==1 for change in changes))

    @patch('bids.services.proposal_final_review.verify_artifact')
    @patch('bids.services.proposal_final_review.build_proposal_pptx')
    def test_regression_or_failure_keeps_original_bytes(self,render,verify):
        file,plan=self.fixture();render.return_value=deepcopy(file)
        edit=lambda slide,inputs:SimpleNamespace(model_dump=lambda:{'slide_changes':[{'slide_number':slide['slide_number'],
            'action':'UPDATE','text_changes':[{'target':slide['elements'][0]['target'],'revised_text':'다른 내용'}]}]})
        self.assertIs(repair_once(file,plan,Mock(),'','',Mock(),Mock(),edit),file)
        self.assertFalse(plan['bounded_repair']['accepted'])
        file,plan=self.fixture();verify.side_effect=TimeoutError()
        self.assertIs(repair_once(file,plan,Mock(),'','',Mock(),Mock(),edit),file)
        self.assertIn('실패',plan['bounded_repair']['reason'])

    def test_verification_failure_prevents_blind_auto_repair(self):
        file,plan=self.fixture();plan['final_document_review']['failures']=['검수 실패'];edit=Mock()
        self.assertIs(repair_once(file,plan,Mock(),'','',Mock(),Mock(),edit),file)
        edit.assert_not_called();self.assertEqual(plan['bounded_repair']['attempt_count'],0)

    def test_uncertain_comparison_alone_never_triggers_automatic_text_rewrite(self):
        file,plan=self.fixture()
        plan['final_document_review']['company_claim_review']['items']=[]
        plan['final_document_review']['conflicts']=[{'first_page':1,'second_page':2,
            'first_quote':'교육 15주','second_quote':'사업 20주','reason':'비교 범위 미확정','verification':'uncertain'}]
        plan['requirement_coverage']['checks']=[]
        edit=Mock()
        self.assertIs(repair_once(file,plan,Mock(),'','',Mock(),Mock(),edit),file)
        edit.assert_not_called()
        self.assertEqual(plan['bounded_repair']['attempt_count'],0)

    @patch('bids.services.proposal_final_review.verify_artifact')
    @patch('bids.services.proposal_final_review.build_proposal_pptx')
    def test_improved_total_cannot_replace_a_previously_confirmed_requirement(self,render,verify):
        file,plan=self.fixture();render.return_value=deepcopy(file)
        plan['requirement_coverage']['checks']=[{'id':'R0001','covered':True},{'id':'R0002','covered':False}]
        plan['output_review']['matched_count']=1
        def verified(content,candidate,*args):
            candidate['final_document_review']['review_required_count']=0
            candidate['requirement_coverage']['checks']=[{'id':'R0001','covered':False},{'id':'R0002','covered':True}]
        verify.side_effect=verified
        edit=lambda slide,inputs:SimpleNamespace(model_dump=lambda:{'slide_changes':[{
            'slide_number':slide['slide_number'],'action':'UPDATE',
            'text_changes':[{'target':slide['elements'][0]['target'],'revised_text':'담당자 확인 필요'}]}]})
        self.assertIs(repair_once(file,plan,Mock(),'','',Mock(),Mock(),edit),file)
        self.assertFalse(plan['bounded_repair']['accepted'])


class FeedbackPipelineTests(SimpleTestCase):
    def test_real_render_is_fully_reviewed_and_original_file_is_preserved(self):
        from pathlib import Path
        from .services.rag.proposal import revise_proposal_with_feedback
        from .services.proposal_output_review import written_pages
        content=deck('기존 첫 페이지','수정하지 않은 회사 실적')
        with TemporaryDirectory() as directory, ExitStack() as stack:
            source=Path(directory)/'draft.pptx';source.write_bytes(content)
            saved=SimpleNamespace(user=Mock(),bid_notice=SimpleNamespace(bid_ntce_no='INTEGRATION'))
            proposal=SimpleNamespace(generated_file=SimpleNamespace(path=str(source)),strategy={},
                                     revision_plan={'requirement_register':{'requirements':[]},
                                         'writing_plan':{'items':[{'id':'R0001','output_slide_numbers':[2],
                                             'detail_slide_numbers':[2]}], 'detail_pages':{'placements':{'R0001':2}}}})
            feedback={'summary':'첫 페이지 수정','slide_changes':[{'slide_number':1,'action':'UPDATE',
                'text_changes':[{'target':'shape-0','revised_text':'수정된 첫 페이지'}]}],
                'added_slides':[],'final_review_items':[]}
            prefix='bids.services.rag.proposal.'
            for name,value in {'search_bid_documents':[], 'build_full_page_context':('',[]),
                'search_web_for_proposal':('',[]),'build_company_knowledge_context':('',{}),
                'company_context':'직접 입력 회사 정보','build_proposal_model':Mock(),'build_text_model':Mock(),
                'build_feedback_plan':SimpleNamespace(model_dump=lambda:feedback)}.items():
                stack.enter_context(patch(prefix+name,return_value=value))
            def chain(prompt,model,schema):
                def invoke(values):
                    if 'blocks' in values:
                        return schema.model_validate({block['id']:{'kind':'company_fact','supported':False,
                            'evidence_quote':'','reason':'확인된 증빙 없음'} for block in json.loads(values['blocks'])})
                    return schema.model_validate({'items':[]})
                return Mock(invoke=Mock(side_effect=invoke))
            stack.enter_context(patch('bids.services.proposal_final_review.structured_chain',side_effect=chain))
            result=revise_proposal_with_feedback(saved,Mock(),proposal,'첫 페이지 변경',slide_number=1)
            self.assertEqual(source.read_bytes(),content)
            self.assertEqual(written_pages(result['file_bytes']),{1:'수정된 첫 페이지',2:'수정하지 않은 회사 실적'})
            self.assertEqual(result['revision_plan']['writing_plan']['items'][0]['output_slide_numbers'],[2])
            self.assertEqual(result['revision_plan']['writing_plan']['detail_pages']['placements'],{'R0001':2})
            self.assertEqual(proposal.revision_plan['writing_plan']['detail_pages'],{'placements':{'R0001':2}})
            report=result['revision_plan']['final_document_review']
            self.assertEqual(report['file_sha256'],sha256(result['file_bytes']).hexdigest())
            self.assertEqual({row['claim'] for row in report['company_claim_review']['items']},
                             {'수정된 첫 페이지','수정하지 않은 회사 실적'})


class FullReviewApiTests(TestCase):
    def setUp(self):
        directory=TemporaryDirectory();self.addCleanup(directory.cleanup)
        override=override_settings(MEDIA_ROOT=directory.name);override.enable();self.addCleanup(override.disable)
        self.user=get_user_model().objects.create_user(username='full-review-owner')
        notice=BidNotice.objects.create(bid_ntce_no='FULL-REVIEW',title='운영 용역',business_type='용역')
        saved=SavedBid.objects.create(user=self.user,bid_notice=notice)
        self.proposal=BidProposal.objects.create(saved_bid=saved,revision_plan={'status':'final'})
        self.content=deck('실제 출력 본문')
        self.proposal.generated_file.save('draft.pptx',ContentFile(self.content))
        self.client=APIClient();self.client.force_authenticate(self.user)

    def post(self):
        return self.client.post('/api/bids/FULL-REVIEW/proposal/full-review/',{},format='json')

    @patch('bids.services.llm.build_text_model',side_effect=AssertionError('No LLM in preflight'))
    def test_submission_check_is_read_only_owner_scoped_and_never_clears_manual_tasks(self,model):
        from pathlib import Path
        path='/api/bids/FULL-REVIEW/proposal/submission-check/'
        self.proposal.revision_plan={'status':'draft','requirement_register':{'requirements':[
            {'requirement':'가격제안서 부가세 포함 제출','sources':['공고 2쪽']}]}}
        self.proposal.save();before=deepcopy(self.proposal.revision_plan);updated=self.proposal.updated_at
        response=self.client.get(path);self.assertEqual(response.status_code,200)
        report=response.data['report']
        self.assertEqual(len(report['manual_checks']),3)
        self.assertEqual(report['separate_requirements'][0]['channel'],'price')
        self.assertTrue(any(row['status']=='unchecked' for row in report['automatic_checks']))
        self.assertEqual(report['file_sha256'],sha256(self.content).hexdigest())
        self.proposal.refresh_from_db()
        self.assertEqual(self.proposal.revision_plan,before);self.assertEqual(self.proposal.updated_at,updated)
        self.assertEqual(Path(self.proposal.generated_file.path).read_bytes(),self.content)
        self.client.force_authenticate(get_user_model().objects.create_user(username='other-preflight'))
        self.assertEqual(self.client.get(path).status_code,404)
        self.client.force_authenticate(None);self.assertEqual(self.client.get(path).status_code,401)

    def test_submission_check_rejects_active_jobs_and_invalid_pptx(self):
        path='/api/bids/FULL-REVIEW/proposal/submission-check/'
        ProposalTask.objects.create(saved_bid=self.proposal.saved_bid,kind='generate')
        self.assertEqual(self.client.get(path).status_code,409)
        ProposalTask.objects.filter(saved_bid=self.proposal.saved_bid).update(status='completed')
        self.proposal.generated_file.save('invalid.pptx',ContentFile(b'broken'))
        self.assertEqual(self.client.get(path).status_code,422)

    @patch('bids.proposal_review_views.build_text_model')
    @patch('bids.proposal_review_views.verify_artifact')
    def test_review_is_owner_scoped_does_not_rewrite_or_finalize_and_uses_no_collection(self,verify,model):
        response=self.post();self.assertEqual(response.status_code,200)
        self.proposal.refresh_from_db()
        self.assertEqual(self.proposal.revision_plan['status'],'final')
        self.assertEqual(self.proposal.generated_file.read(),self.content);self.proposal.generated_file.close()
        self.client.force_authenticate(get_user_model().objects.create_user(username='other-full-review'))
        self.assertEqual(self.post().status_code,404);self.assertEqual(verify.call_count,1)
        self.client.force_authenticate(None);self.assertEqual(self.post().status_code,401)

    @patch('bids.proposal_review_views.build_text_model')
    @patch('bids.proposal_review_views.verify_artifact')
    def test_concurrent_edit_or_external_file_replacement_rejects_old_results(self,verify,model):
        def edit(*args):
            BidProposal.objects.filter(pk=self.proposal.pk).update(updated_at=timezone.now(),revision_plan={'status':'draft','summary':'new'})
        verify.side_effect=edit
        self.assertEqual(self.post().status_code,409)
        self.proposal.refresh_from_db();self.assertEqual(self.proposal.revision_plan['summary'],'new')
        from pathlib import Path
        verify.side_effect=lambda *args:Path(self.proposal.generated_file.path).write_bytes(deck('외부 수정'))
        self.assertEqual(self.post().status_code,409)

    @patch('bids.proposal_review_views.build_text_model')
    def test_generating_and_invalid_packages_do_not_call_a_model(self,model):
        self.proposal.revision_plan={'status':'generating'};self.proposal.save()
        self.assertEqual(self.post().status_code,409);model.assert_not_called()
        self.proposal.revision_plan={'status':'draft'};self.proposal.generated_file.save('bad.pptx',ContentFile(b'broken'))
        self.assertEqual(self.post().status_code,422)

    @patch('bids.services.proposal_tasks.subprocess.Popen')
    def test_async_review_returns_immediately_reuses_task_and_protects_mutations(self,spawn):
        path='/api/bids/FULL-REVIEW/proposal/full-review/'
        response=self.client.post(path,{'async':True},format='json')
        self.assertEqual(response.status_code,202)
        duplicate=self.client.post(path,{'async':True},format='json')
        self.assertEqual(response.data['task']['id'],duplicate.data['task']['id']);spawn.assert_called_once()
        status=self.client.get('/api/bids/FULL-REVIEW/proposal/task/')
        self.assertEqual(status.status_code,200);self.assertIsNone(status.data['proposal'])
        self.assertEqual(self.client.post('/api/bids/FULL-REVIEW/proposal/finalize/',{}).status_code,409)
        self.assertEqual(self.client.post('/api/bids/FULL-REVIEW/proposal/output-review/',{}).status_code,409)
        self.client.force_authenticate(get_user_model().objects.create_user(username='other-job-reader'))
        self.assertEqual(self.client.get('/api/bids/FULL-REVIEW/proposal/task/').status_code,404)

    @patch('bids.services.proposal_tasks.subprocess.Popen')
    @patch('bids.proposal_review_views.review_saved_proposal')
    def test_worker_claims_task_once_and_rejects_prestart_modification(self,review,spawn):
        from rest_framework.response import Response
        review.return_value=Response({},status=200)
        task=enqueue(self.proposal.saved_bid,'review',{})
        run_task(task.pk);run_task(task.pk)
        task.refresh_from_db();self.assertEqual(task.status,'completed');review.assert_called_once()
        task=enqueue(self.proposal.saved_bid,'review',{})
        self.proposal.save()
        run_task(task.pk);task.refresh_from_db()
        self.assertEqual(task.status,'failed');self.assertIn('변경',task.error);self.assertEqual(review.call_count,1)

    @patch('bids.services.proposal_tasks.subprocess.Popen')
    @patch('bids.views.generate_saved_proposal')
    def test_generation_task_delegates_to_validated_generation_helper(self,generate,spawn):
        from rest_framework.response import Response
        from .models import CompanyProfile
        from .services.llm import TASK_KEEP_ALIVE
        CompanyProfile.objects.create(user=self.user,company_name='검증회사',business_registration_number='TEST-001',representative_name='검증')
        def generated(*args):
            self.assertEqual(TASK_KEEP_ALIVE.get(),'1m')
            return Response({},status=201)
        generate.side_effect=generated
        task=enqueue(self.proposal.saved_bid,'generate',{'template_id':'test-template'})
        with patch.dict('os.environ',{'PROPOSAL_LLM_KEEP_ALIVE':'1m'}):
            run_task(task.pk)
        task.refresh_from_db()
        self.assertEqual(task.status,'completed');self.assertEqual(generate.call_args.args[2],'test-template')
        self.assertIsNone(TASK_KEEP_ALIVE.get())

    @patch('bids.services.proposal_tasks.subprocess.Popen')
    @patch('bids.proposal_review_views.review_saved_proposal',side_effect=TimeoutError())
    def test_failed_worker_and_dead_process_restore_status_and_preserve_file(self,review,spawn):
        task=enqueue(self.proposal.saved_bid,'review',{})
        self.proposal.revision_plan={'status':'generating'};self.proposal.save()
        task.payload['proposal_updated_at']=self.proposal.updated_at.isoformat();task.save()
        run_task(task.pk);task.refresh_from_db();self.proposal.refresh_from_db()
        from .services.llm import TASK_KEEP_ALIVE
        self.assertIsNone(TASK_KEEP_ALIVE.get())
        self.assertEqual(task.status,'failed');self.assertEqual(self.proposal.revision_plan['status'],'final')
        task=enqueue(self.proposal.saved_bid,'review',{});task.status='running';task.worker_pid=9999999;task.save()
        self.proposal.revision_plan={'status':'generating'};self.proposal.save()
        with patch('bids.services.proposal_tasks.process_alive',return_value=False):recover_dead_task(task)
        task.refresh_from_db();self.proposal.refresh_from_db()
        self.assertEqual(task.status,'failed');self.assertEqual(self.proposal.revision_plan['status'],'final')
        self.assertEqual(self.proposal.generated_file.read(),self.content);self.proposal.generated_file.close()

    @patch('bids.services.proposal_tasks.subprocess.Popen',side_effect=OSError())
    def test_launch_failure_does_not_leave_active_job(self,spawn):
        with self.assertRaises(ValueError):enqueue(self.proposal.saved_bid,'review',{})
        self.assertFalse(ProposalTask.objects.filter(status__in=['queued','running']).exists())

    def test_process_probe_is_safe_for_current_process(self):
        import os
        self.assertTrue(process_alive(os.getpid()))
