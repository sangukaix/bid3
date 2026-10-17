from hashlib import sha256
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from .models import CompanyDocument
from .services.company_materials import (collect_company_materials, material_context,
    prepare_material_template, apply_company_materials, refresh_material_usage, _asset_file)
from .services.proposal_pptx_renderer import extract_pptx_inventory, get_content_template_numbers
from .services.proposal_submission_check import submission_report


def image_bytes(color='green'):
    buffer=BytesIO();Image.new('RGB',(600,350),color).save(buffer,format='PNG');return buffer.getvalue()


def source_deck(copy='교육 품질 관리 수행 방법 서비스 운영 학습 평가 체계 적용 범위 자료'):
    prs=Presentation()
    cover=prs.slides.add_slide(prs.slide_layouts[6]);cover.shapes.add_textbox(0,0,Inches(8),Inches(1)).text='회사 표지'
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0,0,Inches(9),Inches(1)).text=copy
    group=slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(.6),Inches(1.6),Inches(3),Inches(1)).text='원본 설계 → 준비 → 운영 → 검증'
    slide.shapes.add_picture(BytesIO(image_bytes()),Inches(1),Inches(3),Inches(3))
    slide.shapes.add_picture(BytesIO(image_bytes('blue')),Inches(5),Inches(3),Inches(3))
    slide.shapes.add_picture(BytesIO(image_bytes('red')),Inches(9),Inches(7),Inches(.2))
    buffer=BytesIO();prs.save(buffer);return buffer.getvalue()


class CompanyMaterialTests(TestCase):
    def setUp(self):
        directory=TemporaryDirectory();self.addCleanup(directory.cleanup);self.directory=Path(directory.name)
        override=override_settings(MEDIA_ROOT=directory.name);override.enable();self.addCleanup(override.disable)
        self.user=get_user_model().objects.create_user(username='material-owner')
        self.other=get_user_model().objects.create_user(username='material-other')
        self.original=source_deck()
        self.document=CompanyDocument.objects.create(user=self.user,document_type='proposal',
            original_name='회사 자료.pptx',file=ContentFile(self.original,name='company.pptx'))
        self.template=Path(settings.BASE_DIR)/'proposal_templates/library/learning_sage.pptx'

    def test_owned_catalog_preserves_source_and_labels_groups_and_images(self):
        CompanyDocument.objects.create(user=self.other,document_type='proposal',original_name='외부 회사 비밀.pptx',
            file=ContentFile(source_deck('다른 계정 비밀 내용'),name='foreign.pptx'))
        result=collect_company_materials(self.user)
        self.assertEqual([d['document_id'] for d in result['documents']],[self.document.pk])
        self.assertEqual(len(result['assets']),3)  # 1 diagram region, 2 substantial pictures, no footer logo.
        self.assertEqual(result['assets'][0]['kind'],'source_panel')
        context=material_context(result,'수행 방법 교육 운영')
        self.assertIn('원본 설계 → 준비 → 운영 → 검증',context)
        self.assertIn(f'M{self.document.pk}-P2',context)
        self.assertIn('검토 완료 근거와 구분',context)
        self.assertNotIn('다른 계정 비밀',context)
        self.assertEqual(Path(self.document.file.path).read_bytes(),self.original)
        # A repeat read uses the same content fingerprint and source IDs.
        self.assertEqual(result,collect_company_materials(self.user))

    def test_changed_source_invalidates_catalog_and_old_asset(self):
        old=collect_company_materials(self.user)
        asset=next(a for a in old['assets'] if a['kind']=='image')
        Path(self.document.file.path).write_bytes(source_deck('새 원문 수행 방법 품질 운영'))
        with self.assertRaisesMessage(ValueError,'변경'):_asset_file(asset,self.user)
        new=collect_company_materials(self.user)
        self.assertNotEqual(old['documents'][0]['source_sha256'],new['documents'][0]['source_sha256'])
        self.assertIn('새 원문',new['pages'][1]['text'])

    def test_missing_or_foreign_document_cannot_be_embedded(self):
        result=collect_company_materials(self.user);asset=result['assets'][1]
        with self.assertRaisesMessage(ValueError,'원문'):_asset_file(asset,self.other)
        self.document.delete()
        with self.assertRaisesMessage(ValueError,'원문'):_asset_file(asset,self.user)

    def test_broken_document_is_reported_without_losing_other_sources(self):
        CompanyDocument.objects.create(user=self.user,document_type='company_introduction',
            original_name='broken.pptx',file=ContentFile(b'not a zip',name='broken.pptx'))
        result=collect_company_materials(self.user)
        self.assertEqual(len(result['documents']),1)
        self.assertEqual(result['failed_files'][0]['file_name'],'broken.pptx')

    def test_common_framework_preserves_titles_footer_font_and_slide_count(self):
        materials=collect_company_materials(self.user)
        before=self.template.read_bytes()
        prepared,placements=prepare_material_template(self.template,materials,self.directory/'prepared')
        self.assertTrue(placements)
        source=Presentation(self.template);output=Presentation(prepared)
        self.assertEqual(len(source.slides),len(output.slides))
        self.assertEqual((source.slide_width,source.slide_height),(output.slide_width,output.slide_height))
        for a,b in zip(source.slides,output.slides):
            keep=lambda s:[(x.name,x.text) for x in s.shapes if x.name=='bid3-title' or x.name.startswith('bid3-fixed-')]
            original=keep(a);new=keep(b)
            self.assertTrue(all(item in new for item in original))
        self.assertEqual(self.template.read_bytes(),before)
        inventory=extract_pptx_inventory(prepared)
        available=get_content_template_numbers(inventory)
        self.assertTrue(all(p['source_slide_number'] not in available for p in placements))
        # Neither cover nor table pages are repurposed as source image pages.
        self.assertTrue(all(p['source_slide_number'] not in {1,2,6,7,11,12,15,17,19,20} for p in placements))

    def test_no_relevant_visual_does_not_change_template(self):
        path,items=prepare_material_template(self.template,{'assets':[]},self.directory/'empty')
        self.assertEqual(path,self.template);self.assertEqual(items,[])

    def test_domain_source_slot_keeps_all_specialized_diagrams_and_tables(self):
        materials=collect_company_materials(self.user)
        for asset in materials['assets']:
            asset['text']+=' 현행 서비스 투입 인력 유사 실적 수행 근거 증빙'
        template=Path(settings.PROPOSAL_TEMPLATES['education_learning']['path'])
        prepared,items=prepare_material_template(template,materials,self.directory/'domain',
            bid_title='교육 운영 서비스',select_asset=lambda q,c,t:c[0]['id'])
        self.assertEqual([i['source_slide_number'] for i in items],[13])
        self.assertEqual(items[0]['layout'],'domain-reserved')
        before,after=Presentation(template),Presentation(prepared)
        for index,(a,b) in enumerate(zip(before.slides,after.slides)):
            if index!=12:
                self.assertEqual(a._element.xml,b._element.xml)
            else:
                kept=lambda s:[(x.name,x.text if x.has_text_frame else '') for x in s.shapes
                               if not x.name.startswith('bid3-material-placeholder-') and x.name!='bid3-fixed-material-caption']
                self.assertEqual(kept(a),kept(b))
        self.assertEqual(sum(s.has_table for p in before.slides for s in p.shapes),
                         sum(s.has_table for p in after.slides for s in p.shapes))

    def test_photo_source_caption_replaces_placeholder_notice_without_overlapping_it(self):
        materials=collect_company_materials(self.user)
        template=Path(settings.PROPOSAL_TEMPLATES['education_photo']['path'])
        prepared,items=prepare_material_template(template,materials,self.directory/'photo',
            bid_title='교육 품질 평가 서비스 운영',select_asset=lambda q,c,t:c[0]['id'])
        self.assertTrue(items)
        for item in items:
            page=Presentation(prepared).slides[item['source_slide_number']-1]
            names=[s.name for s in page.shapes]
            self.assertEqual(names.count('bid3-fixed-material-caption'),1)
            self.assertNotIn('bid3-fixed-visual-caption',names)
            text='\n'.join(s.text for s in page.shapes if s.has_text_frame)
            self.assertNotIn('회사 원문 시각 자료를 배치하는 영역',text)
            self.assertIn('현행 적용 범위 확인 필요',text)

    def test_diagram_is_kept_whole_and_source_pages_are_not_repeated(self):
        materials=collect_company_materials(self.user)
        _,items=prepare_material_template(self.template,materials,self.directory/'whole',
            select_asset=lambda section,candidates,title:candidates[0]['id'])
        self.assertEqual(len(items),1)
        self.assertTrue(items[0]['asset_id'].endswith('-PANEL'))

    def test_service_screen_is_not_security_or_performance_evidence(self):
        materials=collect_company_materials(self.user)
        # Score highly on generic business words without providing subject evidence.
        for asset in materials['assets']:
            asset['text']='교육 서비스 수행 운영 신청 학습 관리 화면 품질 검토 적용 범위 자료'
        _,items=prepare_material_template(self.template,materials,self.directory/'screen',
            bid_title='교육 서비스 수행 운영',select_asset=lambda section,candidates,title:
                candidates[0]['id'] if ('정보보호' in section or '수행 역량' in section) else None)
        self.assertEqual(items,[])

    def test_junior_source_panel_is_not_reused_for_employee_training(self):
        materials=collect_company_materials(self.user)
        for asset in materials['assets']:
            asset['text']+=' 주니어 초등 학생 맞춤 교육'
        path,items=prepare_material_template(self.template,materials,self.directory/'audience',
            bid_title='임직원 외국어 교육 운영',select_asset=lambda section,candidates,title:candidates[0]['id'])
        self.assertEqual(path,self.template);self.assertEqual(items,[])

    def test_model_can_reject_candidates_and_failure_keeps_common_template(self):
        materials=collect_company_materials(self.user)
        path,items=prepare_material_template(self.template,materials,self.directory/'rejected',
            select_asset=lambda *args:None)
        self.assertEqual(path,self.template);self.assertEqual(items,[])
        def failed(*args):raise TimeoutError()
        path,items=prepare_material_template(self.template,materials,self.directory/'failed',select_asset=failed)
        self.assertEqual(path,self.template);self.assertEqual(items,[])
        self.assertTrue(materials['warnings'])

    def test_historical_employee_biography_is_not_automatically_reused(self):
        Path(self.document.file.path).write_bytes(source_deck('총괄팀장 주요 진행 프로젝트 경력사항 교육 품질 관리 수행 방법'))
        materials=collect_company_materials(self.user)
        path,items=prepare_material_template(self.template,materials,self.directory/'staff')
        self.assertEqual(path,self.template);self.assertEqual(items,[])

    def test_second_company_uses_same_framework_with_its_own_service_copy(self):
        CompanyDocument.objects.create(user=self.other,document_type='company_introduction',
            original_name='시스템 구축 소개.pptx',file=ContentFile(source_deck(
                '정보시스템 구축 서비스 수행 방법 설계 품질 관리 검증 적용 범위 자료'),'software.pptx'))
        materials=collect_company_materials(self.other)
        self.assertNotIn(self.document.pk,[d['document_id'] for d in materials['documents']])
        prepared,items=prepare_material_template(self.template,materials,self.directory/'software')
        self.assertTrue(items)
        self.assertEqual(len(Presentation(prepared).slides),len(Presentation(self.template).slides))

    def test_picture_crop_is_preserved_and_partial_cache_is_rebuilt(self):
        prs=Presentation(BytesIO(self.original));picture=next(s for s in prs.slides[1].shapes if s.shape_type==13)
        picture.crop_left=.5
        buffer=BytesIO();prs.save(buffer);Path(self.document.file.path).write_bytes(buffer.getvalue())
        materials=collect_company_materials(self.user)
        asset=next(a for a in materials['assets'] if a['kind']=='image')
        with Image.open(Path(asset['_directory'])/asset['file']) as image:self.assertEqual(image.size,(300,350))
        (Path(asset['_directory'])/'catalog.json').write_text('{partial',encoding='utf-8')
        self.assertEqual(collect_company_materials(self.user)['documents'],materials['documents'])

    def test_pdf_visual_regions_are_indexed_but_text_only_pages_are_not_visuals(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        writer=PdfWriter()
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        for index in range(3):
            page=writer.add_blank_page(600,400)
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            stream=DecodedStreamObject();stream.set_data(b'BT /F1 18 Tf 40 300 Td (service delivery process) Tj ET\n'+(b'40 40 100 80 re S\n'*4 if index==1 else b''))
            page[NameObject('/Contents')]=writer._add_object(stream)
        buffer=BytesIO();writer.write(buffer)
        CompanyDocument.objects.create(user=self.other,document_type='proposal',original_name='service.pdf',file=ContentFile(buffer.getvalue(),name='service.pdf'))
        materials=collect_company_materials(self.other)
        self.assertEqual(len(materials['pages']),3)
        self.assertEqual(len(materials['assets']),1)
        self.assertEqual(materials['assets'][0]['page'],2)

    def test_exact_image_is_contained_and_provenance_rebinds_to_actual_pages(self):
        materials=collect_company_materials(self.user)
        image=next(a for a in materials['assets'] if a['kind']=='image')
        materials['assets']=[image]
        prepared,placements=prepare_material_template(self.template,materials,self.directory/'prepared')
        placements=placements[:1];self.assertEqual(len(placements),1)
        number=placements[0]['source_slide_number']
        file={'file_bytes':prepared.read_bytes(),'source_page_map':{i:i for i in range(1,21)},'revision_log':[]}
        result,usage=apply_company_materials(file,materials,placements,self.user)
        prs=Presentation(BytesIO(result['file_bytes']))
        pictures=[s for s in prs.slides[number-1].shapes if s.name.startswith('bid3-company-material-')]
        self.assertEqual(len(pictures),1)
        shape=pictures[0];x,y,w,h=placements[0]['slot']
        self.assertGreaterEqual(shape.left.pt,x-.1);self.assertGreaterEqual(shape.top.pt,y-.1)
        self.assertLessEqual(shape.left.pt+shape.width.pt,x+w+.1)
        self.assertLessEqual(shape.top.pt+shape.height.pt,y+h+.1)
        self.assertEqual(shape.crop_left+shape.crop_right+shape.crop_top+shape.crop_bottom,0)
        self.assertEqual(sha256(shape.image.blob).hexdigest(),usage['used_assets'][0]['image_sha256'])
        self.assertIn('source_sha256',prs.slides[number-1].notes_slide.notes_text_frame.text)
        plan={'company_material_usage':usage};refresh_material_usage(result['file_bytes'],plan)
        self.assertTrue(plan['company_material_usage']['used_assets'][0]['image_matches_source'])
        preflight=submission_report(result['file_bytes'],plan,'')
        check=next(c for c in preflight['automatic_checks'] if c['id']=='company_materials')
        self.assertEqual(check['status'],'attention')
        shape._element.getparent().remove(shape._element);buffer=BytesIO();prs.save(buffer)
        refresh_material_usage(buffer.getvalue(),plan)
        self.assertEqual(plan['company_material_usage']['used_assets'],[])
        self.assertEqual(Path(self.document.file.path).read_bytes(),self.original)

    @patch('bids.services.rag.proposal.select_company_asset',side_effect=lambda section,items,title:items[0]['id'])
    @patch('bids.services.rag.proposal._generate_proposal_from_template')
    def test_generation_passes_original_materials_and_temporary_common_template(self,generate,select):
        from .services.rag.proposal import create_bid_proposal_from_template
        def run(**kwargs):
            self.assertTrue(Path(kwargs['source_path']).exists())
            self.assertTrue(Path(kwargs['material_template_dir']).exists())
            self.assertEqual(kwargs['company_materials']['owner_id'],self.user.pk)
            self.assertTrue(kwargs['company_materials']['assets'])
            return {'checked':True}
        generate.side_effect=run
        result=create_bid_proposal_from_template(SimpleNamespace(user=self.user,bid_notice=SimpleNamespace(title='교육 운영')),None,self.template)
        self.assertTrue(result['checked'])
        self.assertFalse(Path(generate.call_args.kwargs['material_template_dir']).exists())
