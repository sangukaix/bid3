"""Repeatable local-model regression audit of private, manually labelled PPTX cases."""
import json
from io import BytesIO
from hashlib import sha256
from pathlib import Path
from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone
from pptx import Presentation
from pptx.util import Inches, Pt
from bids.services.llm import build_text_model, model_selection
from bids.services.proposal_final_review import verify_artifact
from bids.services.proposal_pptx_renderer import inspect_proposal_quality


class Command(BaseCommand):
    help='Audit a private JSON manifest of source-linked PPTX cases with local models; does not modify DB or source files.'

    def add_arguments(self, parser):
        parser.add_argument('manifest')
        parser.add_argument('--output', required=True)

    def handle(self,*args,**options):
        path=Path(options['manifest']).resolve()
        manifest=json.loads(path.read_text(encoding='utf-8'))
        routes={role:model_selection(role,'unused') for role in ('PROPOSAL','CLAIM_REVIEW')}
        if any(provider!='ollama' for provider,model in routes.values()):
            raise CommandError('회귀 검증은 로컬 모델 전용입니다. 유지보수 페이지에서 두 작업을 Ollama로 설정해 주세요.')
        model=build_text_model('PROPOSAL','unused',3000,reasoning_effort='none')
        reviewer=build_text_model('CLAIM_REVIEW','unused',3000,reasoning_effort='none')
        results=[]
        for case in manifest['cases']:
            self.stdout.write(f"검수 중: {case['name']}")
            if case.get('pptx'):
                content=(path.parent/case['pptx']).read_bytes()
            else:
                presentation=Presentation()
                for text in case['slides']:
                    page=presentation.slides.add_slide(presentation.slide_layouts[6])
                    box=page.shapes.add_textbox(Inches(.6),Inches(.8),Inches(8),Inches(4));box.text=text
                    box.text_frame.paragraphs[0].runs[0].font.size=Pt(22)
                buffer=BytesIO();presentation.save(buffer);content=buffer.getvalue()
            knowledge=(path.parent/case['approved_evidence']).read_text(encoding='utf-8') if case.get('approved_evidence') else ''
            plan={'requirement_register':{'requirements':case.get('requirements',[])}}
            verify_artifact(content,plan,knowledge,model,reviewer)
            quality=inspect_proposal_quality(content)
            checks={row['id']:row for row in plan['requirement_coverage']['checks']}
            errors=[]
            for key,expected in case.get('expected',{}).get('requirements',{}).items():
                if checks.get(key,{}).get('covered') is not expected:
                    errors.append(f'{key}: expected covered={expected}')
            report=plan['final_document_review']
            for expected in case.get('expected',{}).get('claims',[]):
                if not any(expected['text'] in item['claim'] and item['status']==expected['status']
                           for item in report['company_claim_review']['items']):
                    errors.append('회사 주장 판정 불일치: '+expected['text'])
            if 'conflict_count' in case.get('expected',{}) and len(report['conflicts'])!=case['expected']['conflict_count']:
                errors.append('모순 검출 수 불일치')
            if report['failures'] or any(row.get('status')=='unverified' for row in checks.values()):
                errors.append('자동 검수 실패 있음')
            source=case.get('source',{})
            provenance={**source}
            if source.get('file'):
                provenance['file_sha256']=sha256((path.parent/source['file']).read_bytes()).hexdigest()
            results.append({'name':case['name'],'source':provenance,'file_sha256':sha256(content).hexdigest(),
                'expected_checked':bool(case.get('expected')),'expectation_errors':errors,
                'quality_review':quality,**plan})
        report={'version':'proposal-regression-v1','reviewed_at':timezone.now().isoformat(),
            'models':routes,'cases':results,
            'passed':all(case['expected_checked'] and not case['expectation_errors'] for case in results) and bool(results),
            'limitation':'원문·최신 정정공고·실제 회사 증빙 검토 여부는 manifest의 기록입니다. 자동 검수 통과는 제출 적합성 판정이 아닙니다.'}
        output=Path(options['output']).resolve()
        # Never replace a supplied case or manifest with its report.
        inputs={path,*[(path.parent/case['pptx']).resolve() for case in manifest['cases'] if case.get('pptx')]}
        inputs.update((path.parent/case[key]).resolve() for case in manifest['cases'] for key in ('approved_evidence',) if case.get(key))
        inputs.update((path.parent/case['source']['file']).resolve() for case in manifest['cases'] if case.get('source',{}).get('file'))
        if output in inputs:
            raise CommandError('검증 출력은 원문·PPTX·manifest와 다른 파일 경로여야 합니다.')
        output.parent.mkdir(parents=True,exist_ok=True)
        output.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        self.stdout.write(f"완료: {len(results)}개 사례. 기준 일치: {report['passed']}. 결과: {output}")
        if not report['passed']:
            raise CommandError('검증 기준 불일치 또는 수동 기준 미등록: 결과 파일을 확인해 주세요.')
