from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch, MagicMock

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, SimpleTestCase, override_settings
from rest_framework.test import APIClient

from .decks import Deck, inventory, assert_protected, rel_path
from .models import Project, Message, Job
from .workflow import validate_plan, run_job, save_revision, configuration_signature
from . import sources, ai


def example_deck():
    deck=Deck(Path(settings.PROPOSAL_TEMPLATES['learning_sage']['path']).read_bytes())
    deck.order([1,2,3,4,5])
    return deck.bytes()


class DeckTests(SimpleTestCase):
    def test_cloning_copies_chart_and_workbook_instead_of_sharing_them(self):
        from .decks import xml,dump,NS,resolve_part
        from lxml import etree
        deck=Deck(example_deck()); page=deck.slide_parts()[0]
        rels=xml(deck.parts[rel_path(page)])
        etree.SubElement(rels,'{'+NS['rel']+'}Relationship',Id='rIdChart',Type=NS['r']+'/chart',Target='../charts/chart1.xml')
        deck.parts[rel_path(page)]=dump(rels)
        deck.parts['ppt/charts/chart1.xml']=b'<chart/>'
        chart_rels=etree.Element('{'+NS['rel']+'}Relationships')
        etree.SubElement(chart_rels,'{'+NS['rel']+'}Relationship',Id='rIdWorkbook',Type=NS['r']+'/package',Target='../embeddings/data.xlsx')
        deck.parts[rel_path('ppt/charts/chart1.xml')]=dump(chart_rels)
        deck.parts['ppt/embeddings/data.xlsx']=b'original workbook bytes'
        deck.clone(1)
        cloned=deck.slide_parts()[-1]
        copied_chart=resolve_part(cloned,next(r.get('Target') for r in xml(deck.parts[rel_path(cloned)]) if r.get('Id')=='rIdChart'))
        self.assertNotEqual(copied_chart,'ppt/charts/chart1.xml')
        copied_book=resolve_part(copied_chart,xml(deck.parts[rel_path(copied_chart)])[0].get('Target'))
        self.assertNotEqual(copied_book,'ppt/embeddings/data.xlsx')
        self.assertEqual(deck.parts[copied_book],deck.parts['ppt/embeddings/data.xlsx'])

    def test_edit_append_preserves_original_pages_and_native_design(self):
        original=example_deck(); deck=Deck(original); elements=inventory(original)
        number=deck.clone(3); target=elements[2]['elements'][0]['target']
        deck.edit(number,[{'target':target,'text':'학원 프로젝트 발표','font_size':24}])
        changed=deck.bytes()
        assert_protected(original,changed,[1,2,3,4,5])
        self.assertEqual(len(inventory(changed)),6)
        self.assertEqual(inventory(changed)[5]['elements'][0]['text'],'학원 프로젝트 발표')
        old=Deck(original)
        for name in old.parts:
            if name.startswith(('ppt/media/','ppt/slideLayouts/','ppt/theme/')):
                self.assertEqual(old.parts[name],deck.parts[name])

    def test_native_table_cell_dimensions_and_negative_target(self):
        original=Path(settings.PROPOSAL_TEMPLATES['learning_sage']['path']).read_bytes()
        slide=inventory(original)[5]; cell=next(e for e in slide['elements'] if e['kind']=='cell')
        self.assertLess(cell['width'],slide['width']/2)
        deck=Deck(original)
        with self.assertRaises(ValueError): deck.edit(6,[{'target':cell['target'].split(':')[0]+':-1:0','text':'bad'}])
        deck.edit(6,[{'target':cell['target'],'text':'수정된 셀'}])
        self.assertIn('수정된 셀',[e['text'] for e in inventory(deck.bytes())[5]['elements']])

    def test_invalid_pptx_and_deleted_protected_page(self):
        with self.assertRaises(ValueError): Deck(b'not a zip')
        original=example_deck(); shorter=Deck(original); shorter.order([1])
        with self.assertRaisesMessage(ValueError,'잠긴 페이지'): assert_protected(original,shorter.bytes(),[5])

    def test_external_image_rejected(self):
        from .decks import xml, dump, NS
        deck=Deck(example_deck()); path=rel_path(deck.slide_parts()[0]); root=xml(deck.parts[path])
        from lxml import etree
        etree.SubElement(root,'{'+NS['rel']+'}Relationship',Id='rIdUnsafe',Type=NS['r']+'/image',Target='http://127.0.0.1/private',TargetMode='External')
        deck.parts[path]=dump(root)
        with self.assertRaisesMessage(ValueError,'외부 연결'): Deck(deck.bytes())


class SourceTests(SimpleTestCase):
    def test_project_folder_reads_code_but_excludes_secrets_dependencies(self):
        with TemporaryDirectory() as directory:
            root=Path(directory); (root/'app.py').write_text('print("test")')
            (root/'.env').write_text('PASSWORD=do-not-read')
            (root/'secret.json').write_text('private')
            (root/'node_modules').mkdir(); (root/'node_modules'/'index.js').write_text('unwanted')
            text,meta=sources.local_path(str(root))
            self.assertEqual(meta['files'],['app.py']); self.assertIn('print',text)
            self.assertNotIn('do-not-read',text)
            with self.assertRaises(ValueError): sources.local_path(str(root/'.env'))

    def test_public_web_rejects_private_ips_and_redirects(self):
        private=[(2,1,6,'',('127.0.0.1',80))]
        with patch('presentation_studio.sources.socket.getaddrinfo',return_value=private):
            with self.assertRaisesMessage(ValueError,'내부 네트워크'): sources.public_web('http://example.test/')
        public=[(2,1,6,'',('93.184.216.34',80))]
        response=MagicMock(status=302,headers={'Location':'http://127.0.0.1/admin'})
        with patch('presentation_studio.sources.socket.getaddrinfo',side_effect=[public,private]), patch('presentation_studio.sources.urllib3.HTTPConnectionPool') as pool:
            pool.return_value.urlopen.return_value=response
            with self.assertRaises(ValueError): sources.public_web('http://example.test/')
            self.assertEqual(pool.call_args.args[0],'93.184.216.34')
            response.close.assert_called_once()

    def test_html_scripts_are_not_read_as_source(self):
        text,_=sources.extract(b'<h1>Reference</h1><script>bad()</script><p>Content</p>','test.html')
        self.assertIn('Reference',text); self.assertNotIn('bad()',text)


@override_settings(STUDIO_INLINE_JOBS=True)
class StudioTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner=get_user_model().objects.create_user(username='studio-owner',password='only-tests')
        cls.other=get_user_model().objects.create_user(username='studio-other',password='only-tests')
        cls.original=example_deck()

    def setUp(self):
        self.tmp=TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.override=override_settings(MEDIA_ROOT=self.tmp.name); self.override.enable(); self.addCleanup(self.override.disable)
        self.client=APIClient(); self.client.force_authenticate(self.owner)
        result=self.client.post('/api/presentation-studio/projects/',{'title':'학원 프로젝트','file':SimpleUploadedFile('my.pptx',self.original)},format='multipart')
        self.assertEqual(result.status_code,201,result.data)
        self.project=Project.objects.get(pk=result.data['id']); self.root=f'/api/presentation-studio/projects/{self.project.pk}/'
        self.project.template_confirmed=True; self.project.protected_slides=[1,2]; self.project.save()

    def plan(self, **kwargs):
        return {'base_revision':self.project.current_id,'mode':'revise','selected_slide':3,'changes':[],**kwargs}

    def test_authentication_and_owner_scope_all_artifacts(self):
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(self.root).status_code,[401,403])
        self.client.force_authenticate(self.other)
        for path in ['',f'revisions/{self.project.current_id}/download/',f'revisions/{self.project.current_id}/pages/1/','handoff/']:
            self.assertEqual(self.client.get(self.root+path).status_code,404)
        self.assertEqual(self.client.post('/api/presentation-studio/templates/',{'project':str(self.project.pk),'name':'private'},format='json').status_code,404)

    def test_direct_edit_lock_and_optimistic_version(self):
        old=self.project.current_id
        target=self.project.current.inventory[2]['elements'][0]['target']
        body={'base_revision':old,'slide':1,'edits':[{'target':target,'text':'changed'}]}
        self.assertEqual(self.client.post(self.root+'edit/',body,format='json').status_code,400)
        body['slide']=3
        result=self.client.post(self.root+'edit/',body,format='json')
        self.assertEqual(result.status_code,200,result.data)
        self.project.refresh_from_db()
        assert_protected(self.original,Path(self.project.current.file.path).read_bytes(),[1,2])
        self.assertEqual(self.client.post(self.root+'edit/',body,format='json').status_code,400)
        self.assertEqual(self.project.revisions.count(),2)

    def test_reorder_preserves_lock_identity_and_restore_protects(self):
        response=self.client.post(self.root+'edit/',{'action':'order','order':[3,4,5,1,2],'base_revision':self.project.current_id},format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['protected_slides'],[4,5])
        response=self.client.post(self.root+f'revisions/{self.project.current_id}/restore/',{'base_revision':response.data['current']['id']},format='json')
        self.assertEqual(response.status_code,400)

    def test_personal_template_roundtrip_is_owner_scoped(self):
        response=self.client.post('/api/presentation-studio/templates/',{'project':str(self.project.pk),'name':'내 양식'},format='json')
        template=response.data['templates'][0]
        response=self.client.post('/api/presentation-studio/projects/',{'title':'다음 발표','personal_template':template['id']},format='json')
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(response.data['protected_slides'],[1,2])
        copied=Project.objects.get(pk=response.data['id'])
        self.assertEqual(Path(copied.current.file.path).read_bytes(),self.original)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post('/api/presentation-studio/projects/',{'personal_template':template['id']},format='json').status_code,404)

    def test_optional_instruction_and_admin_only_local_paths(self):
        self.assertEqual(self.project.instruction,'')
        self.assertFalse(self.client.get('/api/presentation-studio/capabilities/').data['local_paths_allowed'])
        data={'kind':'path','locator':self.tmp.name}
        self.assertEqual(self.client.post(self.root+'references/',data,format='json').status_code,403)
        self.owner.is_staff=True; self.owner.save()
        (Path(self.tmp.name)/'source.py').write_text('print("new source")')
        with patch('presentation_studio.sources.local_path',return_value=('code',{'files':['source.py']})):
            self.assertEqual(self.client.post(self.root+'references/',data,format='json').status_code,201)

    def test_invalid_reference_reports_error_without_500(self):
        for name in ['broken.pdf','broken.docx','broken.pptx','broken.png']:
            result=self.client.post(self.root+'references/',{'kind':'upload','file':SimpleUploadedFile(name,b'broken')},format='multipart')
            self.assertEqual(result.status_code,400,(name,result.data))

    def test_scope_questions_and_stale_plans_are_blocked(self):
        for plan in [self.plan(questions=['목적은?']),self.plan(base_revision=-1),self.plan(changes=[{'slide':1,'edits':[]}]),
                     self.plan(mode='continue',append_count=1,changes=[{'slide':3,'edits':[]}],additions=[{'template_slide':3}]),
                     self.plan(changes=[{'slide':4,'edits':[]}]),self.plan(mode='discuss',additions=[{'template_slide':3}])]:
            with self.subTest(plan=plan),self.assertRaises(ValueError): validate_plan(self.project,plan)

    def test_apply_appends_and_keeps_all_old_pages(self):
        target=self.project.current.inventory[2]['elements'][0]['target']
        plan=self.plan(mode='continue',append_count=1,additions=[{'template_slide':3,'title':'회고','brief':'제공된 내용만 요약'}])
        msg=Message.objects.create(project=self.project,role='assistant',content='추가합니다',plan=plan)
        with patch('presentation_studio.ai.context',return_value={}),patch('presentation_studio.ai.fill_page',return_value={'edits':[{'target':target,'text':'개발 회고'}],'questions':[]}),patch('presentation_studio.workflow.render'):
            response=self.client.post(self.root+f'plans/{msg.pk}/apply/')
        self.assertEqual(response.status_code,202,response.data)
        self.project.refresh_from_db(); msg.refresh_from_db()
        self.assertEqual(len(self.project.current.inventory),6)
        assert_protected(self.original,Path(self.project.current.file.path).read_bytes(),[1,2,3,4,5])
        self.assertEqual(msg.applied_revision_id,self.project.current_id)
        self.assertEqual(self.project.jobs.latest('created_at').status,'completed')
        self.assertEqual(self.client.post(self.root+f'plans/{msg.pk}/apply/').status_code,400)

    def test_question_during_apply_keeps_original_and_posts_question(self):
        plan=self.plan(mode='continue',append_count=1,additions=[{'template_slide':3,'title':'성과','brief':'성과 작성'}])
        msg=Message.objects.create(project=self.project,role='assistant',content='계획',plan=plan)
        with patch('presentation_studio.ai.context',return_value={}),patch('presentation_studio.ai.fill_page',return_value={'edits':[],'questions':['실제 결과를 알려 주세요.']}):
            self.client.post(self.root+f'plans/{msg.pk}/apply/')
        self.assertEqual(self.project.revisions.count(),1)
        self.assertTrue(self.project.messages.latest('id').plan['questions'])
        self.assertEqual(self.project.jobs.latest('created_at').status,'failed')

    def test_cancelled_worker_cannot_publish_revision(self):
        job=Job.objects.create(project=self.project,kind='apply',status='cancelled')
        with self.assertRaises(ValueError): save_revision(self.project,self.original,'cancelled result',job=job)
        run_job(job.id); self.assertEqual(self.project.revisions.count(),1)

    def test_chat_runs_only_after_design_confirmation_and_keeps_questions(self):
        self.project.template_confirmed=False; self.project.save()
        self.assertEqual(self.client.post(self.root+'chat/',{'message':'만들어줘'},format='json').status_code,400)
        self.project.template_confirmed=True; self.project.save()
        with patch('presentation_studio.ai.plan',return_value={'message':'발표 목적을 정해 볼까요?','questions':['주제는 무엇인가요?']}):
            response=self.client.post(self.root+'chat/',{'message':'만들어줘'},format='json')
        self.assertEqual(response.status_code,202)
        self.assertEqual(self.project.messages.last().plan['questions'],['주제는 무엇인가요?'])

    def test_active_job_blocks_manual_edit_and_settings(self):
        Job.objects.create(project=self.project,kind='chat')
        self.assertEqual(self.client.patch(self.root,{'protected_slides':[]},format='json').status_code,400)
        with self.assertRaises(ValueError): save_revision(self.project,self.original,'concurrent')

    def test_oversized_user_instruction_is_not_silently_truncated(self):
        with patch('presentation_studio.ai.OllamaChatModel') as model:
            with self.assertRaisesMessage(ValueError,'입력 범위'): ai.call(ai.Plan,ai.RULES,{'instruction':'긴 지침 '*16000})
            model.assert_not_called()

    def test_changed_instruction_invalidates_prepared_plan(self):
        plan=self.plan(configuration_signature=configuration_signature(self.project))
        self.project.instruction='새로운 조건'; self.project.save()
        with self.assertRaisesMessage(ValueError,'지침'): validate_plan(self.project,plan)

    def test_image_insertion_uses_owned_reference_and_preserves_locks(self):
        from PIL import Image
        buffer=BytesIO(); Image.new('RGB',(20,20),'blue').save(buffer,format='PNG')
        result=self.client.post(self.root+'references/',{'kind':'upload','file':SimpleUploadedFile('sample.png',buffer.getvalue(),content_type='image/png')},format='multipart')
        ref=result.data['references'][0]
        response=self.client.post(self.root+'edit/',{'action':'image','slide':3,'base_revision':self.project.current_id,'reference':ref['id'],'x':40,'y':100,'width':100},format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.project.refresh_from_db(); content=Path(self.project.current.file.path).read_bytes()
        assert_protected(self.original,content,[1,2])
        from pptx import Presentation
        self.assertTrue(any(s.shape_type==13 for s in Presentation(BytesIO(content)).slides[2].shapes))

    def test_builtin_starter_has_general_presentation_placeholders(self):
        response=self.client.post('/api/presentation-studio/projects/',{'template_id':'learning_sage','initial_pages':3},format='json')
        self.assertEqual(response.status_code,201,response.data)
        text=' '.join(e['text'] for s in response.data['current']['slides'] for e in s['elements'])
        self.assertNotIn('제안',text); self.assertNotIn('발주',text)
        self.assertIn('발표 제목',text)
