"""Surgical OOXML edits. Unchanged slides and their assets keep their original bytes.

python-pptx is used for reading only; round-tripping an uploaded deck through an
object model can discard unsupported artwork, animation and extension data.
"""
from copy import deepcopy
from io import BytesIO
import posixpath
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED, BadZipFile

from lxml import etree as ET
from pptx import Presentation
from PIL import Image

NS = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'rel':'http://schemas.openxmlformats.org/package/2006/relationships',
      'ct':'http://schemas.openxmlformats.org/package/2006/content-types'}
MAX_BYTES = 40 * 1024 * 1024


def xml(data):
    return ET.fromstring(data, ET.XMLParser(resolve_entities=False, no_network=True, load_dtd=False))


def dump(root):
    return ET.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def rel_path(part):
    directory, name = posixpath.split(part)
    return posixpath.join(directory, '_rels', name + '.rels')


def resolve_part(part, target):
    return posixpath.normpath(posixpath.join(posixpath.dirname(part), target)) if not target.startswith('/') else target.lstrip('/')


class Deck:
    def __init__(self, content):
        if len(content) > MAX_BYTES:
            raise ValueError('PPTX는 40MB 이하로 올려 주세요.')
        try:
            with ZipFile(BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > 10000 or sum(i.file_size for i in infos) > 180 * 1024 * 1024:
                    raise ValueError('압축을 푼 PPTX의 크기가 너무 큽니다.')
                if any('..' in i.filename.split('/') or i.filename.startswith('/') or 'vbaproject' in i.filename.lower() for i in infos):
                    raise ValueError('매크로 또는 올바르지 않은 경로가 포함된 파일입니다.')
                self.parts = {i.filename: archive.read(i) for i in infos}
            self.presentation = xml(self.parts['ppt/presentation.xml'])
            self.rels = xml(self.parts['ppt/_rels/presentation.xml.rels'])
            self.types = xml(self.parts['[Content_Types].xml'])
            for name, data in self.parts.items():
                if name.endswith('.rels'):
                    for relation in xml(data):
                        if relation.get('TargetMode') == 'External' and not relation.get('Type','').endswith('/hyperlink'):
                            raise ValueError('외부 연결 그림·개체는 PPTX에 직접 포함한 뒤 올려 주세요.')
            if not 1 <= len(self.slide_parts()) <= 80:
                raise ValueError('PPTX는 1~80페이지를 지원합니다.')
        except (BadZipFile, KeyError, ET.XMLSyntaxError) as error:
            raise ValueError('암호가 없고 정상적으로 열리는 PPTX 파일을 올려 주세요.') from error

    def slide_parts(self):
        relationships = {r.get('Id'):r.get('Target') for r in self.rels}
        return [resolve_part('ppt/presentation.xml', relationships[s.get(f'{{{NS["r"]}}}id')])
                for s in self.presentation.findall('p:sldIdLst/p:sldId', NS)]

    def bytes(self):
        self.parts['ppt/presentation.xml'] = dump(self.presentation)
        self.parts['ppt/_rels/presentation.xml.rels'] = dump(self.rels)
        self.parts['[Content_Types].xml'] = dump(self.types)
        result = BytesIO()
        with ZipFile(result, 'w', ZIP_DEFLATED) as archive:
            for name, content in self.parts.items():
                archive.writestr(name, content)
        return result.getvalue()

    def clone(self, number):
        original = self.slide_parts()
        if not 1 <= number <= len(original) or len(original) >= 80:
            raise ValueError('복제할 페이지 또는 전체 페이지 수를 확인하세요.')
        copied = {}

        def copy_part(old):
            if old in copied:
                return copied[old]
            stem, suffix = posixpath.splitext(old)
            index = 1
            new = f'{stem}_studio{index}{suffix}'
            while new in self.parts:
                index += 1
                new = f'{stem}_studio{index}{suffix}'
            copied[old] = new
            self.parts[new] = self.parts[old]
            for override in list(self.types):
                if override.get('PartName') == '/' + old:
                    item = deepcopy(override); item.set('PartName', '/' + new); self.types.append(item)
                    break
            if rel_path(old) in self.parts:
                relationships = xml(self.parts[rel_path(old)])
                for rel in list(relationships):
                    if rel.get('TargetMode') == 'External':
                        continue
                    kind = rel.get('Type','').rsplit('/',1)[-1]
                    if kind in {'notesSlide','comment','comments'}:
                        relationships.remove(rel)  # Old private notes are not new slide content.
                        continue
                    target = resolve_part(old, rel.get('Target'))
                    # Shared layouts/themes are immutable. Charts, diagrams and their
                    # workbooks are cloned so later edits never change an old slide.
                    mapped = target if kind in {'slideLayout','slideMaster','theme','image'} else copy_part(target)
                    rel.set('Target', posixpath.relpath(mapped, posixpath.dirname(new)))
                self.parts[rel_path(new)] = dump(relationships)
            return new

        part = copy_part(original[number - 1])
        ids = {r.get('Id') for r in self.rels}
        rid = f'rIdStudio{len(ids)+1}'
        while rid in ids: rid += 'x'
        ET.SubElement(self.rels, f'{{{NS["rel"]}}}Relationship', Id=rid,
                      Type=NS['r']+'/slide', Target=posixpath.relpath(part, 'ppt'))
        order = self.presentation.find('p:sldIdLst', NS)
        identifier = max(int(s.get('id')) for s in order) + 1
        ET.SubElement(order, f'{{{NS["p"]}}}sldId', id=str(identifier), attrib={f'{{{NS["r"]}}}id':rid})
        return len(original) + 1

    def order(self, numbers):
        order = self.presentation.find('p:sldIdLst', NS)
        old = list(order)
        if not numbers or len(set(numbers)) != len(numbers) or any(n < 1 or n > len(old) for n in numbers):
            raise ValueError('페이지 순서가 올바르지 않습니다.')
        for element in old: order.remove(element)
        for number in numbers: order.append(old[number - 1])

    def edit(self, number, changes):
        paths = self.slide_parts()
        if not 1 <= number <= len(paths): raise ValueError('페이지를 찾을 수 없습니다.')
        path = paths[number - 1]
        root = xml(self.parts[path])
        for change in changes:
            target = str(change.get('target',''))
            pieces = target.split(':')
            if len(pieces) not in {1,3} or any(not p.isdigit() for p in pieces): raise ValueError('텍스트 위치를 확인하세요.')
            candidates = root.xpath('.//p:cNvPr[@id=$id]', namespaces=NS, id=pieces[0])
            if len(candidates) != 1: raise ValueError('수정할 텍스트 위치를 찾을 수 없습니다.')
            shape = candidates[0].getparent().getparent()
            if len(pieces) == 3:
                table = shape.find('.//a:tbl', NS)
                try: body = table.findall('a:tr',NS)[int(pieces[1])].findall('a:tc',NS)[int(pieces[2])].find('a:txBody',NS)
                except (AttributeError, IndexError, ValueError): raise ValueError('표의 셀 위치가 올바르지 않습니다.')
            else:
                body = shape.find('p:txBody',NS)
            if body is None: raise ValueError('이 개체는 텍스트 편집을 지원하지 않습니다.')
            if 'text' in change:
                text = str(change['text'])
                if len(text) > 5000: raise ValueError('한 텍스트 상자는 5,000자 이하로 입력하세요.')
                old_p = body.find('a:p', NS)
                props = old_p.find('a:pPr', NS) if old_p is not None else None
                run_props = body.find('.//a:rPr', NS)
                end_props = old_p.find('a:endParaRPr',NS) if old_p is not None else None
                for paragraph in list(body.findall('a:p', NS)): body.remove(paragraph)
                for line in text.split('\n'):
                    p = ET.SubElement(body, f'{{{NS["a"]}}}p')
                    if props is not None: p.append(deepcopy(props))
                    run = ET.SubElement(p, f'{{{NS["a"]}}}r')
                    if run_props is not None: run.append(deepcopy(run_props))
                    ET.SubElement(run, f'{{{NS["a"]}}}t').text = line
                    if end_props is not None: p.append(deepcopy(end_props))
            if 'font_size' in change or 'color' in change or 'bold' in change:
                for run in body.findall('.//a:r',NS):
                    props = run.find('a:rPr',NS)
                    if props is None: props = ET.Element(f'{{{NS["a"]}}}rPr'); run.insert(0,props)
                    if 'font_size' in change:
                        size = float(change['font_size'])
                        if not 8 <= size <= 120: raise ValueError('글자 크기는 8~120pt 범위입니다.')
                        props.set('sz', str(round(size * 100)))
                    if 'bold' in change: props.set('b','1' if change['bold'] else '0')
                    if 'color' in change:
                        import re
                        color = str(change['color']).lstrip('#')
                        if not re.fullmatch('[0-9a-fA-F]{6}', color): raise ValueError('색상을 확인하세요.')
                        for fill in props.findall('a:solidFill',NS): props.remove(fill)
                        ET.SubElement(ET.SubElement(props,f'{{{NS["a"]}}}solidFill'),f'{{{NS["a"]}}}srgbClr',val=color)
        self.parts[path] = dump(root)

    def add_image(self, number, image_content, x, y, width):
        with Image.open(BytesIO(image_content)) as image:
            if image.format not in {'PNG','JPEG'} or image.width * image.height > 25_000_000:
                raise ValueError('PNG/JPG 이미지(2,500만 픽셀 이하)를 선택하세요.')
            extension = 'png' if image.format == 'PNG' else 'jpeg'
            height = float(width) * image.height / image.width
        size = self.presentation.find('p:sldSz',NS)
        sw, sh = int(size.get('cx'))/12700, int(size.get('cy'))/12700
        if not (0 <= x < sw and 0 <= y < sh and 5 <= width <= sw-x and height <= sh-y):
            raise ValueError('그림이 페이지 안에 들어오도록 위치와 너비를 줄여 주세요.')
        path = self.slide_parts()[number - 1]
        root = xml(self.parts[path]); tree = root.find('p:cSld/p:spTree',NS)
        index = 1
        while f'ppt/media/studio_{index}.{extension}' in self.parts: index += 1
        asset = f'ppt/media/studio_{index}.{extension}'; self.parts[asset] = image_content
        relationships = xml(self.parts[rel_path(path)]) if rel_path(path) in self.parts else ET.Element(f'{{{NS["rel"]}}}Relationships')
        rid = f'rIdImage{index}'
        while any(r.get('Id') == rid for r in relationships): rid += 'x'
        ET.SubElement(relationships,f'{{{NS["rel"]}}}Relationship',Id=rid,Type=NS['r']+'/image',Target=posixpath.relpath(asset,posixpath.dirname(path)))
        self.parts[rel_path(path)] = dump(relationships)
        if not any(t.get('Extension') == extension for t in self.types):
            ET.SubElement(self.types,f'{{{NS["ct"]}}}Default',Extension=extension,ContentType='image/'+extension)
        shape_id = max([int(n.get('id')) for n in root.findall('.//p:cNvPr',NS)] + [0]) + 1
        pic = ET.SubElement(tree,f'{{{NS["p"]}}}pic')
        nv = ET.SubElement(pic,f'{{{NS["p"]}}}nvPicPr')
        ET.SubElement(nv,f'{{{NS["p"]}}}cNvPr',id=str(shape_id),name='사용자 참고 이미지')
        ET.SubElement(nv,f'{{{NS["p"]}}}cNvPicPr'); ET.SubElement(nv,f'{{{NS["p"]}}}nvPr')
        fill = ET.SubElement(pic,f'{{{NS["p"]}}}blipFill')
        ET.SubElement(fill,f'{{{NS["a"]}}}blip',attrib={f'{{{NS["r"]}}}embed':rid})
        ET.SubElement(ET.SubElement(fill,f'{{{NS["a"]}}}stretch'),f'{{{NS["a"]}}}fillRect')
        sp = ET.SubElement(pic,f'{{{NS["p"]}}}spPr'); transform = ET.SubElement(sp,f'{{{NS["a"]}}}xfrm')
        ET.SubElement(transform,f'{{{NS["a"]}}}off',x=str(round(x*12700)),y=str(round(y*12700)))
        ET.SubElement(transform,f'{{{NS["a"]}}}ext',cx=str(round(width*12700)),cy=str(round(height*12700)))
        ET.SubElement(ET.SubElement(sp,f'{{{NS["a"]}}}prstGeom',prst='rect'),f'{{{NS["a"]}}}avLst')
        self.parts[path] = dump(root)


def inventory(content):
    Deck(content)  # Package limits also apply to library and locally read files.
    prs = Presentation(BytesIO(content))
    result = []
    def elements(shapes):
        found = []
        for shape in shapes:
            if shape.shape_type == 6:
                found.extend(elements(shape.shapes)); continue
            common = {'name':shape.name, 'x':round(shape.left/12700,1),'y':round(shape.top/12700,1),
                      'width':round(shape.width/12700,1),'height':round(shape.height/12700,1)}
            frames = [(str(shape.shape_id), shape.text_frame,'text')] if shape.has_text_frame else []
            if shape.has_table:
                frames = [(f'{shape.shape_id}:{r}:{c}',cell.text_frame,'cell') for r,row in enumerate(shape.table.rows) for c,cell in enumerate(row.cells)]
            for target, frame, kind in frames:
                sizes = [run.font.size.pt for p in frame.paragraphs for run in p.runs if run.font.size]
                bounds=common
                if kind=='cell':
                    _,r,c=map(int,target.split(':'))
                    if shape.table.cell(r,c).is_spanned: continue
                    cell=shape.table.cell(r,c)
                    bounds={**common,'width':round(sum(shape.table.columns[i].width for i in range(c,c+cell.span_width))/12700,1),
                            'height':round(sum(shape.table.rows[i].height for i in range(r,r+cell.span_height))/12700,1)}
                found.append({**bounds,'target':target,'text':frame.text,'kind':kind,'font_size':sizes[0] if sizes else 18})
        return found
    for number, slide in enumerate(prs.slides,1):
        items = elements(slide.shapes)
        title = next((e['text'] for e in items if e['name']=='bid3-title'), next((e['text'] for e in items if e['text'].strip()), f'{number}페이지'))
        result.append({'number':number,'title':title[:150],'elements':items,'width':round(prs.slide_width/12700,1),'height':round(prs.slide_height/12700,1)})
    return result


def assert_protected(before, after, numbers):
    old, new = Deck(before), Deck(after)
    for number in numbers:
        if not 1<=number<=min(len(old.slide_parts()),len(new.slide_parts())): raise ValueError('잠긴 페이지가 사라지는 버전으로 복원할 수 없습니다. 유지 설정을 먼저 확인하세요.')
        a, b = old.slide_parts()[number-1], new.slide_parts()[number-1]
        if old.parts[a] != new.parts[b] or old.parts.get(rel_path(a)) != new.parts.get(rel_path(b)):
            raise ValueError(f'{number}페이지는 유지하도록 잠겨 있습니다.')
