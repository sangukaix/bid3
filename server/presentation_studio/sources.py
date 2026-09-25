"""Read-only, bounded references. No source code or document instruction is executed."""
from html.parser import HTMLParser
from io import BytesIO
import hashlib
import ipaddress
import os
from pathlib import Path
import socket
from urllib.parse import urlparse, urljoin

import urllib3
from docx import Document
from pypdf import PdfReader
from .decks import inventory

TEXT_EXTENSIONS = {'.txt','.md','.csv','.json','.py','.js','.jsx','.ts','.tsx','.html','.css','.sql','.yaml','.yml','.toml','.rst'}
IMAGE_EXTENSIONS = {'.png','.jpg','.jpeg'}
DOCUMENT_EXTENSIONS = {'.pdf','.docx','.pptx'}
EXCLUDED = {'.git','.local','.backups','.codex','.ssh','.aws','node_modules','venv','.venv','__pycache__','.next','media','chroma_db'}
MAX_SOURCE = 16*1024*1024


def safe_name(path):
    name = Path(path).name.lower()
    return not (name.startswith('.') or any(s in name for s in ('credential','secret','password','token','private_key','id_rsa')) or name.endswith(('.key','.pem')))


class TextHTML(HTMLParser):
    def __init__(self): super().__init__(); self.skip = 0; self.text = []
    def handle_starttag(self,tag,attrs):
        if tag in {'script','style','noscript'}: self.skip += 1
        if tag in {'p','div','li','h1','h2','h3','br','tr'}: self.text.append('\n')
    def handle_endtag(self,tag):
        if tag in {'script','style','noscript'}: self.skip = max(0,self.skip-1)
    def handle_data(self,data):
        if not self.skip: self.text.append(data)


def extract(content, filename):
    if len(content) > MAX_SOURCE: raise ValueError('참고 파일은 16MB 이하로 올려 주세요.')
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        try: text = content.decode('utf-8-sig')
        except UnicodeDecodeError: text = content.decode('cp949', errors='replace')
        if suffix == '.html':
            parser = TextHTML(); parser.feed(text); text = ' '.join(parser.text)
    elif suffix == '.pdf':
        document = PdfReader(BytesIO(content))
        if document.is_encrypted: raise ValueError('암호를 해제한 PDF를 올려 주세요.')
        text = '\n\n'.join(f'[페이지 {i+1}]\n{page.extract_text() or ""}' for i,page in enumerate(document.pages[:100]))
    elif suffix == '.docx':
        from zipfile import ZipFile
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80*1024*1024: raise ValueError('문서 압축 크기가 너무 큽니다.')
        doc = Document(BytesIO(content))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    elif suffix == '.pptx':
        text = '\n\n'.join(f'[페이지 {s["number"]}]\n'+ '\n'.join(e['text'] for e in s['elements']) for s in inventory(content))
    elif suffix in IMAGE_EXTENSIONS:
        from PIL import Image
        with Image.open(BytesIO(content)) as image:
            if image.width*image.height > 25_000_000 or image.format not in {'PNG','JPEG'}: raise ValueError('PNG/JPG 2,500만 픽셀 이하 이미지를 올려 주세요.')
            image.verify()
        return '', {'image':True,'vision_pending':True}
    else: raise ValueError('PDF, DOCX, PPTX, PNG/JPG 또는 텍스트·소스 코드 파일을 지원합니다.')
    return text[:200000], {'truncated':len(text)>200000,'empty':not text.strip()}


def public_web(url):
    """Pin each redirect's connection to a validated public IP (no DNS rebinding)."""
    for _ in range(4):
        parsed = urlparse(url)
        if parsed.scheme not in {'http','https'} or not parsed.hostname or parsed.username or parsed.password or parsed.port not in {None,80,443}:
            raise ValueError('공개 HTTP/HTTPS 웹 주소를 입력하세요.')
        host = parsed.hostname.encode('idna').decode('ascii')
        addresses = sorted({item[4][0] for item in socket.getaddrinfo(host,parsed.port or (443 if parsed.scheme=='https' else 80),type=socket.SOCK_STREAM)})
        if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
            raise ValueError('내부 네트워크·로컬 주소는 웹 참고자료로 열 수 없습니다.')
        pool_class = urllib3.HTTPSConnectionPool if parsed.scheme=='https' else urllib3.HTTPConnectionPool
        options = {'assert_hostname':host,'server_hostname':host} if parsed.scheme=='https' else {}
        pool = pool_class(addresses[0],port=parsed.port or (443 if parsed.scheme=='https' else 80),timeout=urllib3.Timeout(connect=5,read=20),**options)
        response=None
        try:
            response = pool.urlopen('GET',(parsed.path or '/') + ('?'+parsed.query if parsed.query else ''),
                                    headers={'Host':host,'User-Agent':'BID3-PresentationStudio/1.0','Accept-Encoding':'identity'},redirect=False,retries=False,preload_content=False)
            if response.status in {301,302,303,307,308}:
                url = urljoin(url,response.headers.get('Location','')); continue
            if response.status != 200: raise ValueError(f'웹 자료를 읽지 못했습니다 (HTTP {response.status}).')
            data = response.read(4*1024*1024+1)
            if len(data)>4*1024*1024: raise ValueError('웹 참고자료 한 페이지는 4MB 이하만 읽습니다.')
            content_type = response.headers.get('Content-Type','').lower()
            if 'pdf' in content_type: text, metadata = extract(data,'web.pdf')
            elif 'html' in content_type or 'text/' in content_type:
                text,metadata = extract(data,'web.html' if 'html' in content_type else 'web.txt')
            else: raise ValueError('웹 링크는 공개 HTML·텍스트·PDF를 지원합니다. 다른 자료는 파일로 올려 주세요.')
            return text, {**metadata,'resolved_url':url}
        finally:
            if response is not None: response.close()
            pool.close()
    raise ValueError('웹 주소의 이동 횟수가 너무 많습니다.')


def local_path(value):
    if not value or value.startswith(('\\\\','//')): raise ValueError('이 서버 PC의 로컬 절대 경로를 입력하세요.')
    requested = Path(value).expanduser()
    if not requested.is_absolute() or not requested.exists(): raise ValueError('서버 PC에서 해당 절대 경로를 찾을 수 없습니다.')
    if any(not safe_name(p) or p.name.lower() in EXCLUDED for p in [requested, *requested.parents] if p.name):
        raise ValueError('숨김·인증정보·실행환경·사용자 데이터 폴더는 읽지 않습니다.')
    root = requested.resolve()
    if root.is_file():
        content = root.read_bytes() if root.stat().st_size <= MAX_SOURCE else b''
        if not content: raise ValueError('빈 파일 또는 크기 제한을 초과한 파일입니다.')
        text,metadata = extract(content,root.name)
        if metadata.get('image'): raise ValueError('이미지는 파일 업로드로 추가해 주세요.')
        return text,{**metadata,'files':[root.name],'resolved_path':str(root)}
    files, parts, total, omitted = [], [], 0, 0
    for directory, folders, names in os.walk(root, followlinks=False):
        relative = Path(directory).relative_to(root)
        folders[:] = sorted(d for d in folders if d.lower() not in EXCLUDED and safe_name(d) and not (Path(directory)/d).is_symlink() and (Path(directory)/d).resolve().is_relative_to(root)) if len(relative.parts)<5 else []
        for name in sorted(names):
            candidate = Path(directory)/name
            if not safe_name(name) or candidate.suffix.lower() not in TEXT_EXTENSIONS: continue
            resolved = candidate.resolve()
            if candidate.is_symlink() or not resolved.is_relative_to(root): continue
            size = resolved.stat().st_size
            if len(files)>=80 or size>400000 or total+size>3_000_000:
                omitted += 1; continue
            text,_ = extract(resolved.read_bytes(),name); total += size
            filename = str(candidate.relative_to(root)); files.append(filename)
            parts.append(f'[파일 {filename}]\n{text}')
    if not files: raise ValueError('읽을 수 있는 텍스트·소스 코드 파일이 없습니다.')
    return '\n\n'.join(parts), {'files':files,'omitted_files':omitted,'truncated':bool(omitted),'resolved_path':str(root),'read_only':True}


def refresh(reference):
    if reference.kind=='url': text,meta = public_web(reference.locator)
    elif reference.kind=='path': text,meta = local_path(reference.locator)
    elif reference.kind=='upload': text,meta = extract(Path(reference.file.path).read_bytes(),reference.name)
    else: text,meta = reference.text,reference.metadata
    reference.text,reference.metadata = text,meta
    reference.digest = hashlib.sha256(text.encode()).hexdigest()
    reference.save(update_fields=['text','metadata','digest','updated_at'])
    return reference
