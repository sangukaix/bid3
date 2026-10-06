from unittest.mock import patch
from django.test import SimpleTestCase, override_settings
from .quality import review_page, semantic_review
from .ai import compact_slide, call, PageText, fill_page
from langchain_core.runnables import RunnableLambda
from bids.services.llm import LocalOutputLimitError
from copy import deepcopy


class QualityTests(SimpleTestCase):
    def setUp(self):
        self.slide={'number':3,'title':'회고','width':960,'height':540,'elements':[
            {'target':'11','text':'원래 제목','kind':'text','width':350,'height':80,'font_size':24},
            {'target':'12','text':'유지할 문구','kind':'text','width':350,'height':80,'font_size':20}]}
        self.packet={'references':[]}

    def test_missing_and_duplicate_targets_are_rejected(self):
        for edits in ([{'target':'not-here','text':'수정'}], [{'target':'11','text':'A'},{'target':'11','text':'B'}]):
            with self.assertRaises(ValueError): review_page(self.slide,edits,'수정',self.packet)

    def test_repair_only_flagged_target_and_recheck(self):
        initial={'findings':[{'target':'11','severity':'error','problem':'근거 없는 개선 수치'}],'citations':[]}
        clean={'findings':[],'citations':[]}
        with patch('presentation_studio.quality.semantic_review',side_effect=[initial,clean]) as review, \
             patch('presentation_studio.quality.call',return_value={'edits':[{'target':'11','text':'측정 필요'}],'questions':[]}) as repair:
            edits,audit=review_page(self.slide,[{'target':'11','text':'90% 개선'},{'target':'12','text':'수업 회고'}],'참고자료만 사용',self.packet)
        self.assertEqual(review.call_count,2); repair.assert_called_once()
        self.assertEqual(edits[1]['text'],'수업 회고')
        self.assertEqual(edits[0]['text'],'측정 필요'); self.assertTrue(audit['repaired_once'])

    def test_second_failure_retains_findings_without_second_repair(self):
        finding={'findings':[{'target':'11','severity':'error','problem':'근거 부족'}],'citations':[]}
        with patch('presentation_studio.quality.semantic_review',return_value=finding), \
             patch('presentation_studio.quality.call',return_value={'edits':[{'target':'11','text':'근거 부족'}],'questions':[]}) as repair:
            with self.assertRaises(ValueError) as raised: review_page(self.slide,[{'target':'11','text':'90% 개선'}],'검사',self.packet)
        repair.assert_called_once(); self.assertTrue(raised.exception.quality_review['findings'])

    def test_repair_cannot_change_other_text(self):
        finding={'findings':[{'target':'11','severity':'error','problem':'문제'}],'citations':[]}
        with patch('presentation_studio.quality.semantic_review',return_value=finding), \
             patch('presentation_studio.quality.call',return_value={'edits':[{'target':'12','text':'범위 밖'}],'questions':[]}):
            with self.assertRaisesMessage(ValueError,'수정 범위'): review_page(self.slide,[{'target':'11','text':'검사'}],'검사',self.packet)

    def test_unknown_citation_is_rejected(self):
        with patch('presentation_studio.quality.call',return_value={'findings':[],'citations':[{'target':'11','reference_id':'fake','part':1}]}):
            with self.assertRaisesMessage(ValueError,'참고자료'): semantic_review(self.slide,[{'target':'11','text':'검사'}],'검사',self.packet)

    def test_citation_to_excluded_source_is_rejected(self):
        packet={'references':[{'id':'real','excerpts':[{'part':1,'text':'실제 원문'}]}]}
        with patch('presentation_studio.quality.call',return_value={'findings':[],
            'citations':[{'target':'11','reference_id':'real','part':1}], '_generation':{'source_ids':[]}}):
            with self.assertRaisesMessage(ValueError,'참고자료'): semantic_review(self.slide,[{'target':'11','text':'검사'}],'검사',packet)

    def test_budget_fit_preserves_packet_and_reports_actual_sources(self):
        packet={'references':[{'id':str(n),'text':'긴 참고자료 '*4000} for n in range(3)],'conversation':[{'content':'옛 대화 '*3000}]}
        before=deepcopy(packet)
        with patch('presentation_studio.ai.OllamaChatModel') as model:
            model.return_value.with_structured_output.return_value.invoke.return_value=PageText(edits=[],questions=[])
            result=call(PageText,'한글로 작성',packet)
        self.assertEqual(packet,before)
        self.assertEqual(result['_generation']['source_ids'],[])
        self.assertEqual(result['_generation']['omitted_reference_count'],3)

    def test_output_limit_retries_only_once(self):
        with patch('presentation_studio.ai.OllamaChatModel') as model:
            invoke=model.return_value.with_structured_output.return_value.invoke
            invoke.side_effect=[LocalOutputLimitError('length'),PageText(edits=[],questions=[])]
            result=call(PageText,'한글로 작성',{})
            self.assertEqual(invoke.call_count,2);self.assertEqual(result['_generation']['retries'],1)
            invoke.reset_mock();invoke.side_effect=LocalOutputLimitError('length')
            with self.assertRaises(LocalOutputLimitError): call(PageText,'한글로 작성',{})
            self.assertEqual(invoke.call_count,2)

    def test_full_text_and_all_targets_reach_the_model(self):
        self.slide['elements'][0]['text']='긴 참고 원문 '*150
        compact=compact_slide(self.slide)
        self.assertEqual(compact['elements'][0]['text'],self.slide['elements'][0]['text'])
        self.assertEqual(compact['elements'][0]['width'],350)

    def test_fixed_template_labels_are_preserved_outside_ai_targets(self):
        self.slide['elements'].append({**self.slide['elements'][0],'target':'footer','name':'bid3-fixed-footer'})
        self.assertEqual({e['target'] for e in compact_slide(self.slide)['elements']},{'11','12'})

    def test_empty_decorations_are_excluded_but_native_blank_text_and_cells_remain(self):
        from io import BytesIO
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        from .decks import inventory
        prs=Presentation(); page=prs.slides.add_slide(prs.slide_layouts[6])
        background=page.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(5),Inches(3))
        blank=page.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1))
        table=page.shapes.add_table(1,1,Inches(1),Inches(3),Inches(2),Inches(1))
        buffer=BytesIO();prs.save(buffer);slide=inventory(buffer.getvalue())[0]
        targets={e['target'] for e in compact_slide(slide)['elements']}
        self.assertNotIn(str(background.shape_id),targets)
        self.assertIn(str(blank.shape_id),targets)
        self.assertIn(f'{table.shape_id}:0:0',targets)
        with self.assertRaisesMessage(ValueError,'배경 도형'):
            review_page(slide,[{'target':str(background.shape_id),'text':'잘못된 배경 글자'}],'작성',{})
        # Saved inventories from before this flag was introduced are conservative.
        for element in slide['elements']: element.pop('ai_editable')
        self.assertNotIn(str(background.shape_id),{e['target'] for e in compact_slide(slide)['elements']})

    def test_native_title_is_used_instead_of_small_section_header(self):
        from io import BytesIO
        from pptx import Presentation
        from pptx.util import Inches
        from .decks import inventory
        prs=Presentation();page=prs.slides.add_slide(prs.slide_layouts[6])
        page.shapes.add_textbox(0,0,Inches(3),Inches(1)).text='1 프로젝트 개요'
        title=page.shapes.add_textbox(0,Inches(1),Inches(5),Inches(1));title.name='title';title.text='특화 포인트'
        buffer=BytesIO();prs.save(buffer)
        self.assertEqual(inventory(buffer.getvalue())[0]['title'],'특화 포인트')

    def test_new_page_missing_title_or_body_retries_once_then_requires_all_targets(self):
        partial={'edits':[{'target':'11','text':'새 제목'}],'questions':[]}
        complete={'edits':[{'target':'11','text':'새 제목'},{'target':'12','text':'새 본문'}],'questions':[]}
        replies=iter([partial,complete])
        with patch('presentation_studio.ai.OllamaChatModel') as model:
            model.return_value.with_structured_output.side_effect=lambda schema: RunnableLambda(lambda messages:schema.model_validate(next(replies)))
            result=fill_page(None,self.slide,'회고',{})
        self.assertEqual(result['_generation']['retries'],1)
        self.assertEqual({e['target'] for e in result['edits']},{'11','12'})
        with patch('presentation_studio.ai.OllamaChatModel') as model:
            model.return_value.with_structured_output.side_effect=lambda schema: RunnableLambda(lambda messages:schema.model_validate(partial))
            with self.assertRaisesMessage(ValueError,'응답의 형식'): fill_page(None,self.slide,'회고',{})
            self.assertEqual(model.return_value.with_structured_output.call_count,2)

    def test_repair_can_ask_user_without_publishing_partial_text(self):
        finding={'findings':[{'target':'11','severity':'error','problem':'수치 근거 필요'}],'citations':[]}
        with patch('presentation_studio.quality.semantic_review',return_value=finding),patch('presentation_studio.quality.call',return_value={'edits':[],'questions':['실제 측정 결과가 있나요?']}):
            with self.assertRaises(ValueError) as raised: review_page(self.slide,[{'target':'11','text':'90% 개선'}],'검사',self.packet)
        self.assertEqual(raised.exception.questions,['실제 측정 결과가 있나요?'])

    @override_settings(STUDIO_SEMANTIC_REVIEW=False)
    def test_layout_failure_requires_repair_and_is_not_silently_clipped(self):
        long='내용을 삭제하지 않고 확인합니다. '*100
        with patch('presentation_studio.quality.call',return_value={'edits':[{'target':'11','text':long}],'questions':[]}) as repair:
            with self.assertRaises(ValueError): review_page(self.slide,[{'target':'11','text':long}],'검사',self.packet)
        repair.assert_called_once()
