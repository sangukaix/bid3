"""Install visually reviewed domain frameworks without changing legacy decks."""
import json,shutil,hashlib,html
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from pptx import Presentation

root=Path(__file__).resolve().parents[2];w=root/'.local/theme-studio-20261014'
review=w/'review';library=root/'server/proposal_templates/library'
configs=json.loads((w/'configs.json').read_text('utf-8'))
catalog=json.loads((library/'catalog.json').read_text('utf-8'))
layouts=json.loads((library/'layouts.json').read_text('utf-8'))
assert len(configs)==18
new_ids={c['id'] for c in configs if not c['sample']}
catalog=[c for c in catalog if c['id'] not in new_ids]
for cfg in configs:
    source=w/cfg.get('output_directory','output-v2')/('samples' if cfg['sample'] else 'templates')
    deck=source/(cfg['stem']+'.pptx');images=source/cfg['stem']
    manifest=json.loads((images/'manifest.json').read_text('utf-8'))
    assert manifest['source_sha256']==hashlib.sha256(deck.read_bytes()).hexdigest()
    assert len(list(images.glob('slide_*.png')))==14
    prs=Presentation(deck)
    assert all('undefined' not in (s.text if s.has_text_frame else '') for p in prs.slides for s in p.shapes)
    destination=review/('samples' if cfg['sample'] else 'templates');destination.mkdir(parents=True,exist_ok=True)
    shutil.copy2(deck,destination/deck.name);shutil.copy2(deck.with_suffix('.pdf'),destination/(cfg['stem']+'.pdf'))
    shutil.copytree(images,destination/cfg['stem'],dirs_exist_ok=True)
    if not cfg['sample']:
        shutil.copy2(deck,library/cfg['file'])
        shutil.copytree(images,library/'previews'/cfg['id'],dirs_exist_ok=True)
        layouts[cfg['file']]=cfg['metadata']
        catalog.append({key:value for key,value in cfg.items() if key not in {'stem','sample','metadata','output_directory'}})
(library/'catalog.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(library/'layouts.json').write_text(json.dumps(layouts,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
sources=json.loads((w/'references/sources.json').read_text('utf-8'))
shutil.copytree(w/'references',review/'references',dirs_exist_ok=True)
(root/'server/proposal_templates/domain_sources.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
explanations={
 'education':'교육 과정과 운영 절차의 설명 순서 참고. 과거 회사·강사·실적을 신규 제안사의 사실로 사용하지 않음.',
 'construction':'현장 제약, 공정·공종별 시공·품질·안전·변경과 인계 구조 참고. 2020년 공고의 기간·계약 조건은 전용하지 않음.',
 'it':'현행·목표 구성, 요구사항 분류, 인터페이스·시험·보안과 관리 산출물의 연결 참고. 2025년 사업 조건은 전용하지 않음.',
 'research':'연구 방법, 자료 출처·신뢰성 검토, 단계별 보고와 인력·보안 관리 구조 참고.',
 'event':'행사 계획, 프로그램·공간·현장 인력, 홍보·참가자 관리·안전과 결과 보고 구조 참고. 2025년 행사의 규모·일정은 전용하지 않음.',
 'supply':'세부 규격, 호환성, 설치·시험·교육·유지보수와 정품·공급 자료의 연결 참고. 2019년 품목·수량은 전용하지 않음.',
 'cmc':'문제 정의, 단계별 계획, 계층·흐름과 검증 근거를 연결하는 시각 설명 방식 참고. 제품·상표·고유 문구는 복사하지 않음.',
 'hasom':'제품 기능·구성·차이점과 운영 조건을 구분하는 설명 방식 참고. 오래된 가격·통신·기술 조건은 전용하지 않음.',
}
rows=['# 분야별 제안서 참고 원문','', '2026-10-14에 공개된 원문 파일을 내려받았습니다. 실제 공개 제안서 3개, 발주처 제안요청서·입찰지침서 5개입니다. 참고 자료의 과거 일정·금액·평가 조건은 현재 입찰조건이 아닙니다. BID3 생성 프롬프트에 개발 지침을 포함하지 않습니다.','']
for s in sources:rows.extend([f"## {s['title']}",f"- 자료 유형: {s['kind']}",f"- 파일: [{s['file']}]({s['file']})",f"- 공개 원문: {s['url']}",f"- 참고한 부분: {explanations[s['id']]}",f"- SHA-256: {s['sha256']}",''])
(review/'references/README.md').write_text('\n'.join(rows),encoding='utf-8')
for name,folder,patterns in [('BID3_templates_12','templates',['*.pptx']),('BID3_samples_6','samples',['*.pptx','*.pdf']),('BID3_reference_originals_8','references',['*.pdf','*.hwp','sources.json','README.md'])]:
    with ZipFile(review/(name+'.zip'),'w',ZIP_DEFLATED) as z:
        for pattern in patterns:
            for file in sorted((review/folder).glob(pattern)):z.write(file,file.name)
samples=[c for c in configs if c['sample']]
viewer=[dict(id=c['theme_id'],name=c['category'],summary=c['description'],stem=c['stem'],variants=[dict(id=t['id'],name=t['name']) for t in configs if not t['sample'] and t['theme_id']==c['theme_id']]) for c in samples]
data=json.dumps(viewer,ensure_ascii=False).replace('<','\\u003c')
refrows=''.join(f'<tr><td>{html.escape(s["kind"])}</td><td><a href="references/{s["file"]}" download>{html.escape(s["title"])}</a></td><td><a href="{html.escape(s["url"],quote=True)}" target="_blank" rel="noreferrer">공개 원문</a></td></tr>' for s in sources)
page='''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>BID3 분야별 제안서 비교</title>
<style>*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#eef1f4;color:#20303c;font-family:'Malgun Gothic',sans-serif}header{padding:45px max(24px,calc((100vw - 1200px)/2));background:#fff}h1{font-size:34px;margin:12px 0}p{line-height:1.8;color:#65717d}.eyebrow{color:#357060;font-size:14px;font-weight:700}.download{display:flex;gap:10px;flex-wrap:wrap}a{color:#275c95}.download a,nav button{border:1px solid #cdd5dc;border-radius:8px;padding:11px 15px;text-decoration:none;background:#fff;font:inherit;font-size:14px;cursor:pointer}nav{display:flex;gap:10px;flex-wrap:wrap;padding:20px max(24px,calc((100vw - 1200px)/2));position:sticky;top:0;z-index:2;background:#eef1f4ee;backdrop-filter:blur(12px)}nav button[aria-pressed=true]{background:#203e39;color:#fff;border-color:#203e39}main{max-width:1200px;margin:0 auto;padding:16px 24px 60px}.intro{background:#fff;padding:25px 30px;border-radius:12px;margin-bottom:30px}.intro h2{margin:0}figure{display:grid;grid-template-columns:45px 1fr;gap:15px;margin:0 0 30px}figure img{width:100%;height:auto;display:block;box-shadow:0 5px 24px #20303c12}figcaption{font-size:17px;font-weight:700;color:#748590;padding-top:15px}section.refs{max-width:1200px;margin:0 auto 50px;padding:28px;background:#fff;border-radius:12px}table{width:100%;border-collapse:collapse;font-size:13px}td{border-bottom:1px solid #e5eaf0;padding:15px;line-height:1.6}td:first-child{width:180px}small{color:#7e8a96}@media(max-width:650px){h1{font-size:26px}nav{position:static;padding:16px}main{padding:16px}figure{grid-template-columns:25px 1fr;gap:6px}td:first-child{width:90px}section.refs{margin:16px;padding:16px}td{padding:9px}}</style>
<header><span class="eyebrow">BID3 PROPOSAL LIBRARY</span><h1>분야마다 다른 제안서 구조</h1><p>6개 분야, 분야별 2개 양식. 각 분야의 대표 샘플 14장을 끝까지 비교해 보세요.<br>교육은 기존 KOTRA 공고 기반 검토용 초안이며, 나머지 5개 분야는 가상 사업의 구조 시연용입니다.</p><div class="download"><a href="BID3_templates_12.zip" download>빈 양식 12종</a><a href="BID3_samples_6.zip" download>샘플 6종 PPTX·PDF</a><a href="BID3_reference_originals_8.zip" download>참고 원문 8개</a><a href="#references">참고 자료 목록</a></div></header>
<nav aria-label="제안서 분야" id="themes"></nav><main><div id="intro" class="intro"></div><div id="slides"></div></main>
<section class="refs" id="references"><h2>내려받은 공개 참고 자료</h2><p>완성 제안서와 발주처의 요청 문서를 구분했습니다. 과거 입찰의 수치·조건을 현재 제안서의 사실로 사용하지 않습니다.</p><table>REFROWS</table><p><a href="references/README.md" download>참고한 부분과 원문 정보</a></p></section>
<script>const themes=DATA;const nav=document.querySelector('#themes');for(const t of themes){const b=document.createElement('button');b.textContent=t.name;b.dataset.id=t.id;b.onclick=()=>show(t.id,true);nav.append(b)}function show(id,scroll=false){const t=themes.find(t=>t.id===id)||themes[0];nav.querySelectorAll('button').forEach(b=>b.setAttribute('aria-pressed',String(b.dataset.id===t.id)));document.querySelector('#intro').innerHTML=`<h2>${t.name}</h2><p>${t.summary}</p><div class="download"><a href="samples/${t.stem}.pptx" download>작성 샘플 PPTX</a><a href="samples/${t.stem}.pdf" download>작성 샘플 PDF</a>${t.variants.map(v=>`<a href="templates/${v.id}.pptx" download>${v.name} 양식</a>`).join('')}</div>`;document.querySelector('#slides').innerHTML=Array.from({length:14},(_,i)=>`<figure><figcaption>${String(i+1).padStart(2,'0')}</figcaption><img src="samples/${t.stem}/slide_${String(i+1).padStart(3,'0')}.png" width="1280" height="720" loading="${i<2?'eager':'lazy'}" alt="${t.name} 제안서 ${i+1}페이지"></figure>`).join('');history.replaceState(null,'','#'+t.id);if(scroll)window.scrollTo({top:180,behavior:'smooth'})}show(location.hash.slice(1));</script></html>'''.replace('REFROWS',refrows).replace('DATA;',data+';')
page=page.replace('</style>','td:last-child{min-width:88px;white-space:nowrap}</style>')
(review/'index.html').write_text(page,encoding='utf-8')
print(json.dumps(dict(installed_templates=12,samples=6,references=len(sources),review=str(review)),ensure_ascii=True))
